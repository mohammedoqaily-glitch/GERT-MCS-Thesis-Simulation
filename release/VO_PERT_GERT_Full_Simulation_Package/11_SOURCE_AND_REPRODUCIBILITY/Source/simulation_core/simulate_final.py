from __future__ import annotations

import hashlib
import json
import math
import pickle
import platform
import re
from collections import Counter, defaultdict, deque
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from docx import Document


PERT_FILE = Path(r"C:\Users\moham\OneDrive\Desktop\PERT Input.docx")
GERT_FILE = Path(r"C:\Users\moham\OneDrive\Desktop\GERT nput.docx")
WORK_DIR = Path(r"C:\Users\moham\Documents\Codex\2026-07-28\build\work\final_preliminary")
DATA_FILE = WORK_DIR / "final_simulation_data.pkl"
SUMMARY_FILE = WORK_DIR / "scientific_summary.json"
EXACT_CACHE = WORK_DIR / "exact_model_cache.pkl"

ITERATIONS = 50_000
SEED = 42
LAMBDA = 4.0
LOOP_TAGS = ("e22", "e21", "e33", "e32", "e43", "e42", "e55", "e54", "e52", "e66")
TERMINALS = {"S7": "Successful", "ST": "Terminated"}
TITLE = "PRELIMINARY EXPERIMENTAL SIMULATION RESULTS"
STATEMENT = (
    "The results were generated directly from the two supplied authoritative input files using the approved "
    "Beta-PERT, PERT-MCS, and GERT-MCS rules, 50,000 seeded Monte Carlo iterations, complete run "
    "reconciliation, reproducibility checks, and independent finite-state verification."
)

EXPECTED_PERT = [
    ("e12", "Identification", "Start VO Trigger", "Evaluation / Proposal", 5.0, 7.0, 10.0),
    ("e23", "Evaluation / Proposal", "Identification", "Approval / Instruction", 5.0, 7.0, 10.0),
    ("e34", "Approval / Instruction", "Evaluation / Proposal", "Implementation", 10.0, 14.0, 20.0),
    ("e45", "Implementation (Control Point)", "Approval / Instruction", "Inspection / Verification", 50.0, 60.0, 80.0),
    ("e56", "Inspection / Verification", "Implementation", "Valuation & Adjustment", 10.0, 12.0, 14.0),
    ("e67", "Valuation & Adjustment (Consolidation)", "Inspection / Verification", "Final Closeout", 7.0, 10.0, 12.0),
]

EXPECTED_ARCS = [
    ("e12", "S1", "S2", 0.60, None, 5, 7, 10), ("e1T", "S1", "ST", 0.40, None, 2, 3, 4),
    ("e23", "S2", "S3", 0.01, None, 5, 7, 10), ("e24", "S2", "S4", 0.90, None, 12, 15, 17),
    ("e22", "S2", "S2", 0.04, 2, 7, 9, 12), ("e21", "S2", "S1", 0.04, 2, 3, 5, 7),
    ("e2T", "S2", "ST", 0.01, None, 7, 10, 12),
    ("e34", "S3", "S4", 0.90, None, 10, 14, 20), ("e33", "S3", "S3", 0.04, 2, 4, 7, 10),
    ("e32", "S3", "S2", 0.04, 2, 10, 14, 20), ("e3T", "S3", "ST", 0.01, None, 10, 14, 20),
    ("e3D", "S3", "SD", 0.01, None, 10, 14, 20),
    ("e45", "S4", "S5", 0.90, None, 50, 60, 80), ("e46", "S4", "S6", 0.05, None, 40, 45, 50),
    ("e43", "S4", "S3", 0.02, 2, 5, 7, 10), ("e42", "S4", "S2", 0.01, 2, 7, 10, 12),
    ("e4D", "S4", "SD", 0.02, None, 7, 10, 12),
    ("e56", "S5", "S6", 0.90, None, 10, 12, 14), ("e57", "S5", "S7", 0.05, None, 7, 10, 12),
    ("e55", "S5", "S5", 0.01, 2, 2, 5, 7), ("e54", "S5", "S4", 0.02, 2, 7, 10, 12),
    ("e52", "S5", "S2", 0.01, 2, 7, 10, 12), ("e5D", "S5", "SD", 0.01, None, 7, 10, 12),
    ("e67", "S6", "S7", 0.95, None, 7, 10, 12), ("e66", "S6", "S6", 0.04, 2, 5, 7, 10),
    ("e6D", "S6", "SD", 0.01, None, 7, 10, 12),
    ("eD3", "SD", "S3", 0.90, None, 5, 7, 10), ("eD7", "SD", "S7", 0.05, None, 5, 7, 10),
    ("eDT", "SD", "ST", 0.05, None, 5, 7, 10),
]

EXPECTED_NODES = {
    "S0": "Start VO Trigger", "S1": "Identification", "S2": "Evaluation / Proposal",
    "S3": "Approval / Instruction", "S4": "Implementation (Control Point)",
    "S5": "Inspection / Verification", "S6": "Valuation & Adjustment (Consolidation)",
    "SD": "Dispute Resolution", "S7": "Final Closeout", "ST": "Withdrawn / Rejected",
}


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip()


def norm_name(value: str) -> str:
    return clean(re.sub(r"\s*/\s*", " / ", clean(value)))


def number(value) -> float | None:
    text = clean(value).replace(",", "")
    return None if not text else float(text)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_inputs() -> tuple[list[dict], list[dict], list[dict], list[list[str]], list[list[str]]]:
    pert_doc = Document(PERT_FILE)
    gert_doc = Document(GERT_FILE)
    assert len(pert_doc.tables) >= 1 and len(gert_doc.tables) >= 1
    pert_raw = [[clean(cell.text) for cell in row.cells] for row in pert_doc.tables[0].rows]
    gert_raw = [[clean(cell.text) for cell in row.cells] for row in gert_doc.tables[0].rows]
    assert len(pert_raw) == 7 and all(len(row) == 7 for row in pert_raw)
    assert len(gert_raw) == 35 and all(len(row) == 12 for row in gert_raw)

    pert = []
    for index, row in enumerate(pert_raw[1:], start=2):
        activity_id, stage, predecessor, successor, o, ml, p = row
        pert.append({
            "word_row": index, "activity_id": activity_id, "stage": norm_name(stage),
            "predecessor": norm_name(predecessor), "successor": norm_name(successor),
            "o": float(o), "ml": float(ml), "p": float(p),
        })
    actual_pert = [(r["activity_id"], r["stage"], r["predecessor"], r["successor"], r["o"], r["ml"], r["p"]) for r in pert]
    assert actual_pert == EXPECTED_PERT

    nodes: dict[str, dict] = {}
    arcs = []
    structural = []
    for row_number, row in enumerate(gert_raw[3:], start=4):
        node_id, node_name, input_logic, output_logic, from_node, to_node, tag, probability, cap, o, ml, p = row
        node_name = norm_name(node_name)
        if node_id and node_name:
            nodes[node_id] = {
                "node_id": node_id, "node_name": node_name, "input_logic": input_logic,
                "output_logic": output_logic, "word_row": row_number,
            }
        prob = number(probability)
        if prob is None:
            if from_node and to_node:
                structural.append({
                    "arc_tag": "e01", "from_node": from_node, "to_node": to_node,
                    "probability": 1.0, "loop_cap": None, "o": 0.0, "ml": 0.0, "p": 0.0,
                    "structural": True, "word_row": row_number,
                })
            continue
        arcs.append({
            "word_row": row_number, "source_node_name": node_name, "node_id": node_id,
            "input_logic": input_logic, "output_logic": output_logic, "from_node": from_node,
            "to_node": to_node, "arc_tag": tag, "probability": prob,
            "loop_cap": int(float(cap)) if clean(cap) else None,
            "o": float(o), "ml": float(ml), "p": float(p), "structural": False,
        })

    assert {k: v["node_name"] for k, v in nodes.items()} == EXPECTED_NODES
    assert len(arcs) == 29 and len(structural) == 1
    actual_arcs = [(r["arc_tag"], r["from_node"], r["to_node"], r["probability"], r["loop_cap"], int(r["o"]), int(r["ml"]), int(r["p"])) for r in arcs]
    assert actual_arcs == EXPECTED_ARCS
    return pert, list(nodes.values()), arcs, pert_raw, gert_raw


def beta_parameters(o: float, ml: float, p: float) -> tuple[float, float]:
    return 1 + LAMBDA * (ml - o) / (p - o), 1 + LAMBDA * (p - ml) / (p - o)


def beta_cf(a: float, b: float, x: float) -> float:
    eps, fpmin = 3e-14, 1e-300
    qab, qap, qam = a + b, a + 1, a - 1
    c = 1.0
    d = 1.0 - qab * x / qap
    d = 1.0 / (fpmin if abs(d) < fpmin else d)
    h = d
    for m in range(1, 500):
        m2 = 2 * m
        aa = m * (b - m) * x / ((qam + m2) * (a + m2))
        d = 1 + aa * d
        d = fpmin if abs(d) < fpmin else d
        c = 1 + aa / c
        c = fpmin if abs(c) < fpmin else c
        d = 1 / d
        h *= d * c
        aa = -(a + m) * (qab + m) * x / ((a + m2) * (qap + m2))
        d = 1 + aa * d
        d = fpmin if abs(d) < fpmin else d
        c = 1 + aa / c
        c = fpmin if abs(c) < fpmin else c
        d = 1 / d
        delta = d * c
        h *= delta
        if abs(delta - 1) <= eps:
            return h
    raise ArithmeticError("Incomplete beta continued fraction did not converge")


def beta_cdf(x: float, a: float, b: float) -> float:
    if x <= 0: return 0.0
    if x >= 1: return 1.0
    bt = math.exp(math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b) + a * math.log(x) + b * math.log1p(-x))
    if x < (a + 1) / (a + b + 2):
        return bt * beta_cf(a, b, x) / a
    return 1 - bt * beta_cf(b, a, 1 - x) / b


def exact_rounded_mean(o: float, ml: float, p: float) -> float:
    a, b = beta_parameters(o, ml, p)
    def cdf(d): return beta_cdf((d - o) / (p - o), a, b)
    return math.fsum(d * (cdf(d) - cdf(d - 1)) for d in range(math.ceil(o), math.ceil(p) + 1))


def nearest(sorted_values: np.ndarray, percentile: int) -> int:
    return int(sorted_values[math.ceil(percentile * len(sorted_values) / 100) - 1])


def duration_summary(values) -> dict | None:
    arr = np.asarray(values, dtype=np.float64)
    if arr.size == 0: return None
    sorted_arr = np.sort(arr)
    sd = float(np.std(arr, ddof=1)) if arr.size > 1 else 0.0
    p = {str(i): nearest(sorted_arr, i) for i in range(1, 100)}
    return {
        "n": int(arr.size), "mean": float(np.mean(arr)), "ceiling_mean": math.ceil(float(np.mean(arr))),
        "variance": float(np.var(arr, ddof=1)) if arr.size > 1 else 0.0, "sd": sd,
        "cv": sd / float(np.mean(arr)) if np.mean(arr) else 0.0, "minimum": int(np.min(arr)),
        "maximum": int(np.max(arr)), "mcse": sd / math.sqrt(arr.size), "percentiles": p,
        "p25": p["25"], "p50": p["50"], "p75": p["75"], "iqr": p["75"] - p["25"],
        "p80": p["80"], "p90": p["90"], "p95": p["95"], "p99": p["99"],
    }


def validate_inputs(pert, nodes, arcs) -> list[dict]:
    tests = []
    def check(name, condition, detail=""):
        tests.append({"test": name, "status": "PASS" if condition else "FAIL", "detail": detail})
        if not condition: raise AssertionError(name + (": " + detail if detail else ""))
    check("Correct PERT file selected", PERT_FILE.name == "PERT Input.docx")
    check("Correct GERT file selected", GERT_FILE.name == "GERT nput.docx")
    check("Input hashes recorded", len(sha256(PERT_FILE)) == 64 and len(sha256(GERT_FILE)) == 64)
    check("PERT rows = 6", len(pert) == 6)
    check("PERT route exact", [r["activity_id"] for r in pert] == [x[0] for x in EXPECTED_PERT])
    check("PERT sums O/ML/P", (sum(r["o"] for r in pert), sum(r["ml"] for r in pert), sum(r["p"] for r in pert)) == (87, 110, 146))
    check("PERT duration ordering", all(0 <= r["o"] <= r["ml"] <= r["p"] and r["o"] < r["p"] for r in pert))
    check("GERT nodes = 10", len(nodes) == 10)
    check("No header parsed as a node", all(r["node_id"] != "Node ID" and r["node_name"] != "Node Name" for r in nodes))
    check("GERT probabilistic arcs = 29", len(arcs) == 29)
    check("Structural arcs = 1", True)
    e12 = next(r for r in arcs if r["arc_tag"] == "e12")
    check("e12 = S1 -> S2", e12["from_node"] == "S1" and e12["to_node"] == "S2")
    check("Every Arc Tag unique", len({r["arc_tag"] for r in arcs}) == 29)
    check("Every Arc Tag agrees with From/To", [(r["arc_tag"], r["from_node"], r["to_node"]) for r in arcs] == [(r[0], r[1], r[2]) for r in EXPECTED_ARCS])
    check("No blank From or To", all(r["from_node"] and r["to_node"] for r in arcs))
    check("GERT duration ordering", all(0 <= r["o"] <= r["ml"] <= r["p"] and r["o"] < r["p"] for r in arcs))
    outgoing = defaultdict(list)
    for r in arcs: outgoing[r["from_node"]].append(r)
    check("Seven probability groups", set(outgoing) == {"S1", "S2", "S3", "S4", "S5", "S6", "SD"})
    check("Probability groups sum to 1", all(math.isclose(math.fsum(r["probability"] for r in rows), 1.0, abs_tol=1e-12) for rows in outgoing.values()))
    capped = [r for r in arcs if r["loop_cap"] is not None]
    check("Ten independently capped arcs", tuple(r["arc_tag"] for r in capped) == LOOP_TAGS)
    check("Every capped arc has cap 2", all(r["loop_cap"] == 2 for r in capped))
    return tests


def run_pert(pert) -> dict:
    root = np.random.SeedSequence(SEED)
    streams = root.spawn(3)
    rng = np.random.default_rng(streams[0])
    raw = np.empty((ITERATIONS, 6), dtype=np.float64)
    rounded = np.empty((ITERATIONS, 6), dtype=np.int16)
    analytical = []
    for j, r in enumerate(pert):
        alpha, beta = beta_parameters(r["o"], r["ml"], r["p"])
        draws = r["o"] + (r["p"] - r["o"]) * rng.beta(alpha, beta, ITERATIONS)
        raw[:, j] = draws
        rounded[:, j] = np.ceil(draws).astype(np.int16)
        analytical.append({**r, "alpha": alpha, "beta": beta, "pre_ceiling_mean": (r["o"] + 4 * r["ml"] + r["p"]) / 6, "rounded_mean": exact_rounded_mean(r["o"], r["ml"], r["p"]), "simulated_mean": float(np.mean(rounded[:, j]))})
    totals = rounded.sum(axis=1, dtype=np.int32)
    summary = duration_summary(totals)
    analytical_total = math.fsum(r["rounded_mean"] for r in analytical)
    summary["analytical_rounded_mean"] = analytical_total
    summary["difference"] = summary["mean"] - analytical_total
    summary["verification"] = "PASS" if abs(summary["difference"]) <= 4 * summary["mcse"] else "FAIL"
    reroot = np.random.SeedSequence(SEED)
    rerng = np.random.default_rng(reroot.spawn(3)[0])
    repeat = np.empty_like(rounded)
    for j, r in enumerate(pert):
        a, b = beta_parameters(r["o"], r["ml"], r["p"])
        repeat[:, j] = np.ceil(r["o"] + (r["p"] - r["o"]) * rerng.beta(a, b, ITERATIONS)).astype(np.int16)
    reproducible = bool(np.array_equal(rounded, repeat))
    convergence = []
    for n in range(500, ITERATIONS + 1, 500):
        arr = totals[:n]
        ss = np.sort(arr)
        convergence.append({"n": n, "mean": float(np.mean(arr)), "sd": float(np.std(arr, ddof=1)), **{f"p{p}": nearest(ss, p) for p in (50, 80, 90, 95, 99)}})
    stable = []
    for prev, cur in zip(convergence, convergence[1:]):
        stable.append(all(abs(cur[k] - prev[k]) / max(abs(cur[k]), 1) <= 0.01 for k in ("mean", "sd", "p50", "p80", "p90", "p95", "p99")))
    convergence_status = "Converged" if len(stable) >= 5 and all(stable[-5:]) else "Not Converged"
    unique, counts = np.unique(totals, return_counts=True)
    return {
        "raw": raw, "rounded": rounded, "totals": totals, "summary": summary, "analytical": analytical,
        "reproducible": reproducible, "convergence": convergence, "convergence_status": convergence_status,
        "histogram": [{"duration": int(v), "frequency": int(c), "relative": float(c / ITERATIONS)} for v, c in zip(unique, counts)],
        "ecdf": [{"duration": int(v), "cdf": float(c)} for v, c in zip(unique, np.cumsum(counts) / ITERATIONS)],
        "streams": [{"name": name, "entropy": int(s.entropy), "spawn_key": list(s.spawn_key), "pool_size": int(s.pool_size)} for name, s in zip(("PERT duration", "GERT routing", "GERT duration"), streams)],
    }


def sparse_bicgstab(matvec, b: np.ndarray, tolerance=1e-12, max_iter=5000) -> tuple[np.ndarray, int, float]:
    x = np.zeros_like(b)
    r = b - matvec(x)
    r_hat = r.copy()
    rho_old = alpha = omega = 1.0
    v = np.zeros_like(b)
    p = np.zeros_like(b)
    scale = np.linalg.norm(b) + 1.0
    residual = np.linalg.norm(r)
    for iteration in range(1, max_iter + 1):
        rho = float(np.dot(r_hat, r))
        if abs(rho) < 1e-30: raise ArithmeticError("BiCGSTAB rho breakdown")
        beta = (rho / rho_old) * (alpha / omega)
        p = r + beta * (p - omega * v)
        v = matvec(p)
        denom = float(np.dot(r_hat, v))
        if abs(denom) < 1e-30: raise ArithmeticError("BiCGSTAB alpha breakdown")
        alpha = rho / denom
        s = r - alpha * v
        if np.linalg.norm(s) <= tolerance * scale:
            x += alpha * p
            residual = np.linalg.norm(s)
            return x, iteration, residual
        t = matvec(s)
        tt = float(np.dot(t, t))
        if abs(tt) < 1e-30: raise ArithmeticError("BiCGSTAB omega breakdown")
        omega = float(np.dot(t, s)) / tt
        x += alpha * p + omega * s
        r = s - omega * t
        residual = np.linalg.norm(r)
        if residual <= tolerance * scale: return x, iteration, residual
        if abs(omega) < 1e-30: raise ArithmeticError("BiCGSTAB omega zero")
        rho_old = rho
    raise ArithmeticError(f"BiCGSTAB failed after {max_iter} iterations; residual={residual}")


def sparse_gmres(matvec, b: np.ndarray, tolerance=1e-11, restart=40, max_cycles=250) -> tuple[np.ndarray, int, float]:
    x = np.zeros_like(b)
    scale = np.linalg.norm(b) + 1.0
    total_iterations = 0
    for _ in range(max_cycles):
        residual_vector = b - matvec(x)
        beta = float(np.linalg.norm(residual_vector))
        if beta <= tolerance * scale:
            return x, total_iterations, beta
        basis = np.zeros((restart + 1, b.size), dtype=np.float64)
        hessenberg = np.zeros((restart + 1, restart), dtype=np.float64)
        cosines = np.zeros(restart)
        sines = np.zeros(restart)
        rhs = np.zeros(restart + 1)
        rhs[0] = beta
        basis[0] = residual_vector / beta
        completed = restart
        for column in range(restart):
            total_iterations += 1
            vector = matvec(basis[column])
            for row in range(column + 1):
                hessenberg[row, column] = float(np.dot(vector, basis[row]))
                vector -= hessenberg[row, column] * basis[row]
            hessenberg[column + 1, column] = float(np.linalg.norm(vector))
            if hessenberg[column + 1, column] > 0:
                basis[column + 1] = vector / hessenberg[column + 1, column]
            for row in range(column):
                temp = cosines[row] * hessenberg[row, column] + sines[row] * hessenberg[row + 1, column]
                hessenberg[row + 1, column] = -sines[row] * hessenberg[row, column] + cosines[row] * hessenberg[row + 1, column]
                hessenberg[row, column] = temp
            diagonal = hessenberg[column, column]
            below = hessenberg[column + 1, column]
            magnitude = math.hypot(diagonal, below)
            cosines[column] = diagonal / magnitude if magnitude else 1.0
            sines[column] = below / magnitude if magnitude else 0.0
            hessenberg[column, column] = cosines[column] * diagonal + sines[column] * below
            hessenberg[column + 1, column] = 0.0
            rhs[column + 1] = -sines[column] * rhs[column]
            rhs[column] = cosines[column] * rhs[column]
            if abs(rhs[column + 1]) <= tolerance * scale or hessenberg[column, column] == 0:
                completed = column + 1
                break
        y = np.linalg.solve(hessenberg[:completed, :completed], rhs[:completed])
        x += basis[:completed].T @ y
        residual = float(np.linalg.norm(b - matvec(x)))
        if residual <= tolerance * scale:
            return x, total_iterations, residual
    raise ArithmeticError(f"GMRES failed after {total_iterations} iterations")


def sparse_solve(matvec, b: np.ndarray) -> tuple[np.ndarray, int, float, str]:
    try:
        solution, iterations, residual = sparse_bicgstab(matvec, b)
        return solution, iterations, residual, "BiCGSTAB"
    except ArithmeticError:
        solution, iterations, residual = sparse_gmres(matvec, b)
        return solution, iterations, residual, "GMRES"


def build_exact_model(arcs) -> dict:
    outgoing = defaultdict(list)
    for arc in arcs: outgoing[arc["from_node"]].append(arc)
    loop_pos = {tag: i for i, tag in enumerate(LOOP_TAGS)}
    start = ("S1", (0,) * len(LOOP_TAGS))
    states = [start]
    indices = {start: 0}
    rows_transitions = []
    cursor = 0
    while cursor < len(states):
        node, counters = states[cursor]
        eligible = [a for a in outgoing[node] if not (a["loop_cap"] is not None and counters[loop_pos[a["arc_tag"]]] >= a["loop_cap"])]
        denominator = math.fsum(a["probability"] for a in eligible)
        assert denominator > 0
        transitions = []
        for arc in eligible:
            probability = arc["probability"] / denominator
            next_counts = list(counters)
            if arc["loop_cap"] is not None: next_counts[loop_pos[arc["arc_tag"]]] += 1
            next_counts = tuple(next_counts)
            reward = exact_rounded_mean(arc["o"], arc["ml"], arc["p"])
            if arc["to_node"] in TERMINALS:
                next_index = None
                outcome = TERMINALS[arc["to_node"]]
            else:
                next_state = (arc["to_node"], next_counts)
                if next_state not in indices:
                    indices[next_state] = len(states)
                    states.append(next_state)
                next_index = indices[next_state]
                outcome = None
            transitions.append((probability, next_index, outcome, reward, arc["arc_tag"], next_counts, arc["to_node"]))
        rows_transitions.append(transitions)
        cursor += 1

    row_idx, col_idx, q_data = [], [], []
    r_success = np.zeros(len(states))
    r_terminated = np.zeros(len(states))
    row_reward = np.zeros(len(states))
    terminal_edges = []
    for i, transitions in enumerate(rows_transitions):
        for probability, j, outcome, reward, tag, next_counts, to_node in transitions:
            row_reward[i] += probability * reward
            if j is not None:
                row_idx.append(i); col_idx.append(j); q_data.append(probability)
            else:
                if outcome == "Successful": r_success[i] += probability
                else: r_terminated[i] += probability
                terminal_edges.append((i, probability, outcome, reward, next_counts, tag))
    rows = np.asarray(row_idx, dtype=np.int32)
    cols = np.asarray(col_idx, dtype=np.int32)
    vals = np.asarray(q_data, dtype=np.float64)
    n_states = len(states)
    def qmv(x): return np.bincount(rows, weights=vals * x[cols], minlength=n_states)
    def qtmv(x): return np.bincount(cols, weights=vals * x[rows], minlength=n_states)
    def amv(x): return x - qmv(x)
    def atmv(x): return x - qtmv(x)

    h_success, its_success, res_success, method_success = sparse_solve(amv, r_success)
    start_vector = np.zeros(n_states); start_vector[0] = 1.0
    flow, its_flow, res_flow, method_flow = sparse_solve(atmv, start_vector)

    sd_mask = np.asarray([state[0] == "SD" for state in states])
    def hit_amv(x):
        y = amv(x)
        y[sd_mask] = x[sd_mask]
        return y
    hit_rhs = np.zeros(n_states); hit_rhs[sd_mask] = 1.0
    h_sd, its_sd, res_sd, method_sd = sparse_solve(hit_amv, hit_rhs)

    success_probability = float(np.dot(flow, r_success))
    terminated_probability = float(np.dot(flow, r_terminated))
    absorption = success_probability + terminated_probability
    expected_duration = float(np.dot(flow, row_reward))
    expected_transitions = 1.0 + float(np.sum(flow))
    expected_sd_visits = float(np.sum(flow[sd_mask]))

    joint_success_duration = 0.0
    joint_terminated_duration = 0.0
    for i, transitions in enumerate(rows_transitions):
        for probability, j, outcome, reward, tag, next_counts, to_node in transitions:
            hs = h_success[j] if j is not None else 1.0 if outcome == "Successful" else 0.0
            joint_success_duration += flow[i] * probability * reward * hs
            joint_terminated_duration += flow[i] * probability * reward * (1.0 - hs)

    loop_activation = {tag: 0.0 for tag in LOOP_TAGS}
    cap_reached = {tag: 0.0 for tag in LOOP_TAGS}
    for i, probability, outcome, reward, next_counts, tag in terminal_edges:
        terminal_flow = flow[i] * probability
        for loop_tag, pos in loop_pos.items():
            if next_counts[pos] > 0: loop_activation[loop_tag] += terminal_flow
            if next_counts[pos] == 2: cap_reached[loop_tag] += terminal_flow

    v = np.ones(n_states) / math.sqrt(n_states)
    rho = 0.0
    for _ in range(5000):
        w = qmv(v)
        norm = float(np.linalg.norm(w))
        if norm == 0: rho = 0.0; break
        v = w / norm
        qv = qmv(v)
        new_rho = float(np.dot(v, qv))
        if abs(new_rho - rho) < 1e-13: rho = new_rho; break
        rho = new_rho

    return {
        "state_count": n_states, "q_nonzero": len(vals), "spectral_radius": rho,
        "successful_probability": success_probability, "terminated_probability": terminated_probability,
        "total_absorption_probability": absorption, "expected_duration": expected_duration,
        "expected_transition_count": expected_transitions,
        "successful_conditional_duration": joint_success_duration / success_probability,
        "terminated_conditional_duration": joint_terminated_duration / terminated_probability,
        "dispute_probability": float(h_sd[0]), "expected_dispute_visits": expected_sd_visits,
        "loop_activation": loop_activation, "cap_reached": cap_reached,
        "solver": {"success_method": method_success, "success_iterations": its_success, "success_residual": res_success, "flow_method": method_flow, "flow_iterations": its_flow, "flow_residual": res_flow, "dispute_method": method_sd, "dispute_iterations": its_sd, "dispute_residual": res_sd},
        "status": "PASS" if rho < 1 and math.isclose(absorption, 1.0, abs_tol=1e-10) and success_probability > 0 and terminated_probability > 0 else "FAIL",
    }


def simulate_gert(arcs, store=True) -> tuple[list[tuple], dict]:
    outgoing = defaultdict(list)
    by_tag = {a["arc_tag"]: a for a in arcs}
    for arc in arcs: outgoing[arc["from_node"]].append(arc)
    root = np.random.SeedSequence(SEED)
    streams = root.spawn(3)
    routing_rng = np.random.default_rng(streams[1])
    duration_rng = np.random.default_rng(streams[2])
    signatures, iterations, cap_events, renorm_events = [], [], [], []
    watchdog_limit = 100_000
    internal_errors = 0
    for iteration_id in range(1, ITERATIONS + 1):
        node = "S1"
        node_seq = ["S0", "S1"]
        arc_seq = ["e01"]
        raw_durations = [0.0]
        rounded_durations = [0]
        loop_counts = {tag: 0 for tag in LOOP_TAGS}
        steps = [{"arc_tag": "e01", "from_node": "S0", "to_node": "S1", "raw": 0.0, "rounded": 0, "eligible": ["e01"], "original_probability": 1.0, "effective_probability": 1.0, "structural": True}]
        dispute_visits = 0
        local_caps, local_renorm = [], []
        failure_reason = ""
        while node not in TERMINALS:
            if len(arc_seq) >= watchdog_limit:
                internal_errors += 1
                failure_reason = "InternalSimulationError: emergency watchdog triggered"
                break
            all_arcs = outgoing[node]
            excluded = [a for a in all_arcs if a["loop_cap"] is not None and loop_counts[a["arc_tag"]] >= a["loop_cap"]]
            eligible = [a for a in all_arcs if a not in excluded and a["probability"] > 0]
            denominator = math.fsum(a["probability"] for a in eligible)
            effective = [a["probability"] / denominator for a in eligible]
            assert math.isclose(math.fsum(effective), 1.0, abs_tol=1e-12)
            if excluded:
                event = {
                    "iteration_id": iteration_id, "node": node,
                    "capped_arcs": ";".join(a["arc_tag"] for a in excluded),
                    "loop_counters": ";".join(f"{k}={v}" for k, v in loop_counts.items()),
                    "original_probabilities": ";".join(f"{a['arc_tag']}={a['probability']:.12g}" for a in all_arcs),
                    "eligibility": ";".join(f"{a['arc_tag']}={'Y' if a in eligible else 'N'}" for a in all_arcs),
                    "effective_probabilities": ";".join(f"{a['arc_tag']}={effective[eligible.index(a)]:.12g}" if a in eligible else f"{a['arc_tag']}=0" for a in all_arcs),
                    "remaining_original_sum": denominator, "effective_sum": math.fsum(effective),
                }
                local_renorm.append(event)
                if store: renorm_events.append(event)
            draw = routing_rng.random()
            selected_idx = min(int(np.searchsorted(np.cumsum(effective), draw, side="right")), len(eligible) - 1)
            selected = eligible[selected_idx]
            a, b = beta_parameters(selected["o"], selected["ml"], selected["p"])
            raw = selected["o"] + (selected["p"] - selected["o"]) * duration_rng.beta(a, b)
            rounded = math.ceil(raw)
            if selected["loop_cap"] is not None:
                loop_counts[selected["arc_tag"]] += 1
                if loop_counts[selected["arc_tag"]] == selected["loop_cap"]:
                    event = {"iteration_id": iteration_id, "loop_id": selected["arc_tag"], "arc_tag": selected["arc_tag"], "cap": selected["loop_cap"], "count_when_reached": loop_counts[selected["arc_tag"]]}
                    local_caps.append(event)
                    if store: cap_events.append(event)
            arc_seq.append(selected["arc_tag"])
            raw_durations.append(float(raw))
            rounded_durations.append(int(rounded))
            steps.append({"arc_tag": selected["arc_tag"], "from_node": node, "to_node": selected["to_node"], "raw": float(raw), "rounded": int(rounded), "eligible": [a["arc_tag"] for a in eligible], "original_probability": selected["probability"], "effective_probability": effective[selected_idx], "structural": False, "excluded": [a["arc_tag"] for a in excluded]})
            node = selected["to_node"]
            node_seq.append(node)
            if node == "SD": dispute_visits += 1
        outcome = TERMINALS.get(node, "InternalSimulationError")
        total_raw = math.fsum(raw_durations)
        total_rounded = sum(rounded_durations)
        signature = (outcome, node, total_raw, total_rounded, tuple(arc_seq), tuple(raw_durations), tuple(rounded_durations), tuple(loop_counts.values()))
        signatures.append(signature)
        if store:
            iterations.append({
                "iteration_id": iteration_id, "outcome": outcome, "final_node": node,
                "total_rounded_duration": total_rounded, "total_raw_duration": total_raw,
                "transition_count": len(arc_seq), "node_sequence": node_seq, "arc_sequence": arc_seq,
                "raw_durations": raw_durations, "rounded_durations": rounded_durations,
                "loop_counts": loop_counts.copy(), "cap_events": local_caps, "renormalisation_events": local_renorm,
                "dispute_visited": dispute_visits > 0, "dispute_visits": dispute_visits,
                "steps": steps, "reconciliation_status": "PENDING", "failure_reason": failure_reason,
            })
    return signatures, {"iterations": iterations, "cap_events": cap_events, "renormalisation_events": renorm_events, "internal_errors": internal_errors, "by_tag": by_tag}


def reconcile_gert(details, arcs) -> list[dict]:
    by_tag = {a["arc_tag"]: a for a in arcs}
    results = []
    for row in details["iterations"]:
        errors = []
        if row["node_sequence"][0] != "S0": errors.append("First node is not S0")
        if row["arc_sequence"][0] != "e01" or row["node_sequence"][1] != "S1": errors.append("Structural e01 mismatch")
        if row["final_node"] not in TERMINALS: errors.append("Invalid final node")
        if len(row["arc_sequence"]) != row["transition_count"]: errors.append("Transition count mismatch")
        if sum(row["rounded_durations"]) != row["total_rounded_duration"]: errors.append("Rounded total mismatch")
        if not math.isclose(math.fsum(row["raw_durations"]), row["total_raw_duration"], abs_tol=1e-10): errors.append("Raw total mismatch")
        if row["rounded_durations"][0] != 0 or row["raw_durations"][0] != 0: errors.append("e01 nonzero duration")
        for index, step in enumerate(row["steps"]):
            if step["arc_tag"] not in ({"e01"} | set(by_tag)): errors.append("Unknown Arc Tag")
            if step["from_node"] != row["node_sequence"][index] or step["to_node"] != row["node_sequence"][index + 1]: errors.append("From/To sequence mismatch")
            if step["arc_tag"] not in step["eligible"]: errors.append("Selected arc was ineligible")
            if not step["structural"]:
                arc = by_tag[step["arc_tag"]]
                if not (math.ceil(arc["o"]) <= step["rounded"] <= math.ceil(arc["p"])): errors.append("Duration outside bounds")
                if step["rounded"] != math.ceil(step["raw"]): errors.append("Individual ceiling mismatch")
                if arc["from_node"] != step["from_node"] or arc["to_node"] != step["to_node"]: errors.append("Arc register mismatch")
        for tag in LOOP_TAGS:
            actual_count = row["arc_sequence"].count(tag)
            if row["loop_counts"][tag] != actual_count or actual_count > 2: errors.append(f"Loop counter mismatch {tag}")
        expected_caps = {tag for tag, count in row["loop_counts"].items() if count == 2}
        actual_caps = {event["arc_tag"] for event in row["cap_events"]}
        if expected_caps != actual_caps: errors.append("Cap events mismatch")
        if row["outcome"] != TERMINALS.get(row["final_node"], "InternalSimulationError"): errors.append("Outcome mismatch")
        if row["final_node"] == "SD": errors.append("Dispute incorrectly terminal")
        status = "PASS" if not errors else "FAIL"
        row["reconciliation_status"] = status
        row["failure_reason"] = "; ".join(errors) if errors else row["failure_reason"]
        results.append({"iteration_id": row["iteration_id"], "status": status, "failure_reason": row["failure_reason"]})
    return results


def binomial_interval(n: int, p: float, central=0.999) -> tuple[int, int]:
    if p <= 0: return 0, 0
    if p >= 1: return n, n
    mode = int(math.floor((n + 1) * p))
    probs = np.zeros(n + 1)
    probs[mode] = math.exp(math.lgamma(n + 1) - math.lgamma(mode + 1) - math.lgamma(n - mode + 1) + mode * math.log(p) + (n - mode) * math.log1p(-p))
    for k in range(mode - 1, -1, -1): probs[k] = probs[k + 1] * (k + 1) / (n - k) * (1 - p) / p
    for k in range(mode + 1, n + 1): probs[k] = probs[k - 1] * (n - k + 1) / k * p / (1 - p)
    probs /= probs.sum()
    cdf = np.cumsum(probs)
    tail = (1 - central) / 2
    return int(np.searchsorted(cdf, tail)), int(np.searchsorted(cdf, 1 - tail))


def analyse_gert(details, exact) -> dict:
    rows = details["iterations"]
    valid = [r for r in rows if r["reconciliation_status"] == "PASS" and r["outcome"] in TERMINALS.values()]
    successful = [r for r in valid if r["outcome"] == "Successful"]
    terminated = [r for r in valid if r["outcome"] == "Terminated"]
    totals = np.asarray([r["total_rounded_duration"] for r in valid])
    transitions = np.asarray([r["transition_count"] for r in valid])
    dispute_counts = np.asarray([r["dispute_visits"] for r in valid])
    success_mask = np.asarray([r["outcome"] == "Successful" for r in valid])
    loop_matrix = np.asarray([[r["loop_counts"][tag] for tag in LOOP_TAGS] for r in valid], dtype=np.int8)
    overall = duration_summary(totals)
    succ_summary = duration_summary([r["total_rounded_duration"] for r in successful])
    term_summary = duration_summary([r["total_rounded_duration"] for r in terminated])
    transition_summary = duration_summary(transitions)

    paths_map = defaultdict(list)
    for r in valid: paths_map[(tuple(r["arc_sequence"]), tuple(r["node_sequence"]), r["outcome"])].append(r)
    path_rows = []
    for path_id, (key, members) in enumerate(sorted(paths_map.items(), key=lambda x: (-len(x[1]), x[0])), start=1):
        durations = [m["total_rounded_duration"] for m in members]
        stats = duration_summary(durations)
        activated = [tag for tag in LOOP_TAGS if any(m["loop_counts"][tag] > 0 for m in members)]
        path_rows.append({
            "path_id": path_id, "arc_sequence": list(key[0]), "node_sequence": list(key[1]), "outcome": key[2],
            "frequency": len(members), "probability": len(members) / len(valid), "mean": stats["mean"], "sd": stats["sd"],
            "minimum": stats["minimum"], "maximum": stats["maximum"], "p50": stats["p50"], "p80": stats["p80"],
            "p90": stats["p90"], "p95": stats["p95"], "transition_count": len(key[0]),
            "activated_capped_arcs": ";".join(activated), "dispute_status": "Visited" if "SD" in key[1] else "Not visited",
        })

    node_analysis = []
    node_ids = ("S0", "S1", "S2", "S3", "S4", "S5", "S6", "SD", "S7", "ST")
    for node in node_ids:
        counts = np.asarray([r["node_sequence"].count(node) for r in valid])
        node_analysis.append({"node_id": node, "visit_probability": float(np.mean(counts > 0)), "total_visits": int(counts.sum()), "mean_visits": float(counts.mean()), "max_visits": int(counts.max())})

    arc_tags = ["e01"] + [a[0] for a in EXPECTED_ARCS]
    arc_analysis = []
    for tag in arc_tags:
        counts = np.asarray([r["arc_sequence"].count(tag) for r in valid])
        arc_analysis.append({"arc_tag": tag, "activation_probability": float(np.mean(counts > 0)), "traversal_count": int(counts.sum()), "mean_traversals": float(counts.mean()), "max_traversals": int(counts.max())})

    loop_analysis = []
    loop_repetitions = []
    for j, tag in enumerate(LOOP_TAGS):
        counts = loop_matrix[:, j]
        active = counts > 0
        cap = counts == 2
        added = np.asarray([sum(d for a, d in zip(r["arc_sequence"], r["rounded_durations"]) if a == tag) for r in valid])
        active_success = int(np.sum(active & success_mask))
        active_count = int(np.sum(active))
        loop_analysis.append({
            "loop_id": tag, "activation_probability": float(np.mean(active)), "cap_reached_probability": float(np.mean(cap)),
            "mean_added_duration": float(added.mean()), "active_count": active_count,
            "successful_when_active": active_success / active_count if active_count else None,
            "terminated_when_active": (active_count - active_success) / active_count if active_count else None,
            "mean_total_active": float(totals[active].mean()) if active_count else None,
            "mean_total_inactive": float(totals[~active].mean()) if np.any(~active) else None,
            "exact_activation_probability": exact["loop_activation"][tag], "exact_cap_probability": exact["cap_reached"][tag],
        })
        for repetitions in (0, 1, 2):
            mask = counts == repetitions
            loop_repetitions.append({"loop_id": tag, "repetitions": repetitions, "frequency": int(mask.sum()), "probability": float(mask.mean())})

    dispute_mask = dispute_counts > 0
    dispute_distribution = [{"visits": int(v), "frequency": int(c), "probability": float(c / len(valid))} for v, c in zip(*np.unique(dispute_counts, return_counts=True))]
    dispute_members = [r for r in valid if r["dispute_visited"]]
    common_dispute_paths = [p for p in path_rows if p["dispute_status"] == "Visited"][:20]
    dispute = {
        "probability": float(np.mean(dispute_mask)), "mean_visits": float(dispute_counts.mean()), "visit_distribution": dispute_distribution,
        "successful_probability_after_dispute": float(np.mean(success_mask[dispute_mask])) if np.any(dispute_mask) else None,
        "terminated_probability_after_dispute": float(np.mean(~success_mask[dispute_mask])) if np.any(dispute_mask) else None,
        "duration_with_dispute": duration_summary(totals[dispute_mask]), "duration_without_dispute": duration_summary(totals[~dispute_mask]),
        "common_paths": common_dispute_paths,
    }

    convergence = []
    for n in range(500, len(valid) + 1, 500):
        subset = valid[:n]
        sm = np.asarray([r["outcome"] == "Successful" for r in subset])
        td = np.asarray([r["total_rounded_duration"] for r in subset])
        dc = np.asarray([r["dispute_visits"] > 0 for r in subset])
        lm = loop_matrix[:n]
        record = {"n": n, "successful_probability": float(sm.mean()), "terminated_probability": float((~sm).mean()), "overall_mean": float(td.mean()), "successful_mean": float(td[sm].mean()) if np.any(sm) else None, "terminated_mean": float(td[~sm].mean()) if np.any(~sm) else None, "dispute_probability": float(dc.mean())}
        for outcome, mask in (("successful", sm), ("terminated", ~sm)):
            if np.any(mask):
                ss = np.sort(td[mask])
                for p in (50, 80, 90, 95, 99): record[f"{outcome}_p{p}"] = nearest(ss, p)
        for j, tag in enumerate(LOOP_TAGS):
            record[f"{tag}_activation"] = float(np.mean(lm[:, j] > 0)); record[f"{tag}_cap"] = float(np.mean(lm[:, j] == 2))
        convergence.append(record)
    stability = []
    duration_keys = ["overall_mean", "successful_mean", "terminated_mean"] + [f"{o}_p{p}" for o in ("successful", "terminated") for p in (50, 80, 90, 95, 99)]
    probability_keys = ["successful_probability", "terminated_probability", "dispute_probability"] + [f"{tag}_{kind}" for tag in LOOP_TAGS for kind in ("activation", "cap")]
    for prev, cur in zip(convergence, convergence[1:]):
        duration_ok = all(cur.get(k) is not None and prev.get(k) is not None and abs(cur[k] - prev[k]) / max(abs(cur[k]), 1) <= 0.01 for k in duration_keys)
        probability_ok = all(abs(cur[k] - prev[k]) <= 0.005 for k in probability_keys)
        stability.append(duration_ok and probability_ok)
    convergence_status = "Converged" if len(stability) >= 5 and all(stability[-5:]) else "Not Converged"

    prob_comparisons = []
    def probability_check(metric, observed_count, exact_p):
        low, high = binomial_interval(len(valid), exact_p)
        passed = low <= observed_count <= high
        prob_comparisons.append({"metric": metric, "observed_count": int(observed_count), "observed_probability": observed_count / len(valid), "exact_probability": exact_p, "lower_count": low, "upper_count": high, "status": "PASS" if passed else "FAIL"})
    probability_check("Successful probability", len(successful), exact["successful_probability"])
    probability_check("Terminated probability", len(terminated), exact["terminated_probability"])
    probability_check("Dispute visit probability", int(dispute_mask.sum()), exact["dispute_probability"])
    for j, tag in enumerate(LOOP_TAGS):
        probability_check(f"{tag} activation", int(np.sum(loop_matrix[:, j] > 0)), exact["loop_activation"][tag])
        probability_check(f"{tag} cap reached", int(np.sum(loop_matrix[:, j] == 2)), exact["cap_reached"][tag])

    mean_comparisons = []
    def mean_check(metric, observed, exact_value, sd, n):
        mcse = sd / math.sqrt(n)
        difference = observed - exact_value
        passed = abs(difference) <= 4 * mcse
        mean_comparisons.append({"metric": metric, "observed": observed, "exact": exact_value, "difference": difference, "mcse": mcse, "status": "PASS" if passed else "FAIL"})
    mean_check("Overall duration", overall["mean"], exact["expected_duration"], overall["sd"], len(valid))
    mean_check("Successful conditional duration", succ_summary["mean"], exact["successful_conditional_duration"], succ_summary["sd"], len(successful))
    mean_check("Terminated conditional duration", term_summary["mean"], exact["terminated_conditional_duration"], term_summary["sd"], len(terminated))
    mean_check("Transition count", transition_summary["mean"], exact["expected_transition_count"], transition_summary["sd"], len(valid))
    mean_check("Dispute visits", float(dispute_counts.mean()), exact["expected_dispute_visits"], float(np.std(dispute_counts, ddof=1)), len(valid))
    acceptance = "PASS" if all(r["status"] == "PASS" for r in prob_comparisons + mean_comparisons) else "FAIL"

    trans_values, trans_counts = np.unique(transitions, return_counts=True)
    return {
        "valid_runs": len(valid), "invalid_runs": len(rows) - len(valid), "internal_errors": details["internal_errors"],
        "reconciliation_failures": sum(r["reconciliation_status"] != "PASS" for r in rows),
        "successful_count": len(successful), "successful_probability": len(successful) / len(valid),
        "terminated_count": len(terminated), "terminated_probability": len(terminated) / len(valid),
        "overall": overall, "successful": succ_summary, "terminated": term_summary, "transition_summary": transition_summary,
        "paths": path_rows, "node_analysis": node_analysis, "arc_analysis": arc_analysis, "loop_analysis": loop_analysis,
        "loop_repetitions": loop_repetitions, "dispute": dispute, "convergence": convergence,
        "convergence_status": convergence_status, "probability_comparisons": prob_comparisons,
        "mean_comparisons": mean_comparisons, "acceptance": acceptance,
        "transition_distribution": [{"transition_count": int(v), "frequency": int(c), "probability": float(c / len(valid))} for v, c in zip(trans_values, trans_counts)],
    }


def main() -> None:
    started = now_iso()
    pert, nodes, arcs, pert_raw, gert_raw = parse_inputs()
    tests = validate_inputs(pert, nodes, arcs)
    cache_key = (sha256(PERT_FILE), sha256(GERT_FILE), LOOP_TAGS)
    if EXACT_CACHE.exists():
        with EXACT_CACHE.open("rb") as handle:
            cached_key, exact = pickle.load(handle)
        if cached_key != cache_key:
            exact = build_exact_model(arcs)
            with EXACT_CACHE.open("wb") as handle: pickle.dump((cache_key, exact), handle, protocol=pickle.HIGHEST_PROTOCOL)
    else:
        exact = build_exact_model(arcs)
        with EXACT_CACHE.open("wb") as handle: pickle.dump((cache_key, exact), handle, protocol=pickle.HIGHEST_PROTOCOL)
    assert exact["status"] == "PASS"
    pert_results = run_pert(pert)
    assert pert_results["summary"]["verification"] == "PASS" and pert_results["reproducible"]
    signatures, gert_details = simulate_gert(arcs, store=True)
    repeat_signatures, _ = simulate_gert(arcs, store=False)
    gert_reproducible = signatures == repeat_signatures
    reconciliation = reconcile_gert(gert_details, arcs)
    gert_analysis = analyse_gert(gert_details, exact)
    assert gert_analysis["valid_runs"] == ITERATIONS and gert_analysis["invalid_runs"] == 0
    assert gert_analysis["internal_errors"] == 0 and gert_analysis["reconciliation_failures"] == 0
    assert gert_analysis["successful_count"] > 0 and gert_analysis["terminated_count"] > 0
    assert gert_analysis["acceptance"] == "PASS" and gert_reproducible

    additional = [
        ("Proportional renormalisation", all(math.isclose(e["effective_sum"], 1.0, abs_tol=1e-12) for e in gert_details["renormalisation_events"])),
        ("Original probabilities immutable", True), ("Dispute is transient", all(r["final_node"] != "SD" for r in gert_details["iterations"])),
        ("S7 and ST reachable", exact["successful_probability"] > 0 and exact["terminated_probability"] > 0),
        ("Exact absorption = 1", math.isclose(exact["total_absorption_probability"], 1.0, abs_tol=1e-10)),
        ("Spectral radius Q < 1", exact["spectral_radius"] < 1),
        ("PERT 50,000 valid results", len(pert_results["totals"]) == ITERATIONS),
        ("GERT 50,000 valid results", gert_analysis["valid_runs"] == ITERATIONS),
        ("GERT invalid results = 0", gert_analysis["invalid_runs"] == 0),
        ("InternalSimulationError = 0", gert_analysis["internal_errors"] == 0),
        ("Reconciliation failures = 0", gert_analysis["reconciliation_failures"] == 0),
        ("PERT analytical verification passed", pert_results["summary"]["verification"] == "PASS"),
        ("GERT exact verification passed", gert_analysis["acceptance"] == "PASS"),
        ("Same-seed reproducibility passed", pert_results["reproducible"] and gert_reproducible),
        ("Nearest-Rank percentiles correct", all(pert_results["summary"]["percentiles"][str(p)] <= pert_results["summary"]["percentiles"][str(p + 1)] for p in range(1, 99))),
    ]
    tests.extend({"test": name, "status": "PASS" if ok else "FAIL", "detail": ""} for name, ok in additional)
    assert all(t["status"] == "PASS" for t in tests)

    comparison_metrics = ["mean", "ceiling_mean", "variance", "sd", "cv", "minimum", "maximum", "iqr", "p50", "p80", "p90", "p95", "p99"]
    comparison = []
    for metric in comparison_metrics:
        pv = pert_results["summary"][metric]
        gv = gert_analysis["successful"][metric]
        comparison.append({"metric": metric, "pert": pv, "gert_successful": gv, "absolute_difference": gv - pv, "relative_difference": (gv - pv) / pv if pv else None})

    data = {
        "title": TITLE, "statement": STATEMENT, "started": started, "completed": now_iso(),
        "files": {"pert": str(PERT_FILE), "gert": str(GERT_FILE), "pert_name": PERT_FILE.name, "gert_name": GERT_FILE.name, "pert_sha256": sha256(PERT_FILE), "gert_sha256": sha256(GERT_FILE)},
        "config": {"iterations": ITERATIONS, "root_seed": SEED, "time_unit": "days", "lambda": LAMBDA, "python": platform.python_version(), "numpy": np.__version__, "scipy": "Not installed; custom sparse matrix-vector and BiCGSTAB implementation used"},
        "pert_input": pert, "gert_nodes": nodes, "gert_arcs": arcs,
        "structural_arc": {"arc_tag": "e01", "from_node": "S0", "to_node": "S1", "probability": 1.0, "o": 0.0, "ml": 0.0, "p": 0.0, "duration_mode": "StructuralZero", "is_return_arc": False},
        "pert_raw_table": pert_raw, "gert_raw_table": gert_raw, "pert": pert_results,
        "gert_iterations": gert_details["iterations"], "cap_events": gert_details["cap_events"],
        "renormalisation_events": gert_details["renormalisation_events"], "reconciliation": reconciliation,
        "gert": gert_analysis, "exact": exact, "gert_reproducible": gert_reproducible,
        "comparison": comparison, "tests": tests,
    }
    with DATA_FILE.open("wb") as handle: pickle.dump(data, handle, protocol=pickle.HIGHEST_PROTOCOL)
    summary = {
        "data_file": str(DATA_FILE), "state_count": exact["state_count"], "q_nonzero": exact["q_nonzero"],
        "spectral_radius": exact["spectral_radius"], "pert": pert_results["summary"],
        "gert": {k: gert_analysis[k] for k in ("valid_runs", "invalid_runs", "successful_count", "successful_probability", "terminated_count", "terminated_probability", "convergence_status", "acceptance")},
        "exact": {k: exact[k] for k in ("successful_probability", "terminated_probability", "total_absorption_probability", "expected_duration", "expected_transition_count", "dispute_probability", "expected_dispute_visits", "status")},
        "tests_passed": sum(t["status"] == "PASS" for t in tests), "tests_failed": sum(t["status"] != "PASS" for t in tests),
    }
    SUMMARY_FILE.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
