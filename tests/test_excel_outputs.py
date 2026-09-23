import csv
import json
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parents[1]
OUTPUT = BASE / "outputs" / "Sensitivity_Analysis"


class ExcelOutputTests(unittest.TestCase):
    def test_required_workbook_sheets_exist(self):
        path = OUTPUT / "Sensitivity_and_Robustness_Results.xlsx"
        self.assertTrue(path.exists() and zipfile.is_zipfile(path))
        required = [
            "README", "Baseline_Config", "Baseline_Validation", "Sensitivity_Parameters", "Routing_OAT",
            "Duration_OAT", "LoopCap_Robustness", "Distribution_Robustness", "Seed_Robustness",
            "Replication_Convergence", "Threshold_Analysis", "CI_Summary", "Tornado_P90_Data",
            "Tornado_Outcome_Data", "Integrated_Ranking", "Scenario_Run_Log", "Validation_Checks", "Errors_Warnings",
        ]
        workbook = load_workbook(path, read_only=True, data_only=False)
        self.assertTrue(set(required).issubset(workbook.sheetnames))
        workbook.close()

    def test_thesis_workbook_has_exactly_seven_tables(self):
        path = OUTPUT / "Thesis_Ready_Tables.xlsx"
        self.assertTrue(path.exists() and zipfile.is_zipfile(path))
        workbook = load_workbook(path, read_only=True, data_only=False)
        self.assertEqual(len(workbook.sheetnames), 7)
        self.assertTrue(all(name.startswith("Table_4") for name in workbook.sheetnames))
        workbook.close()

    def test_main_workbook_embeds_all_publication_figures(self):
        path = OUTPUT / "Sensitivity_and_Robustness_Results.xlsx"
        with zipfile.ZipFile(path) as archive:
            media = [name for name in archive.namelist() if name.startswith("xl/media/")]
        self.assertEqual(len(media), 12)

    def test_all_required_figures_exist_in_both_formats(self):
        figure_dir = OUTPUT / "Figures"
        self.assertEqual(len(list(figure_dir.glob("*.svg"))), 12)
        self.assertEqual(len(list(figure_dir.glob("*.png"))), 12)
        self.assertTrue(all(path.stat().st_size > 0 for path in figure_dir.glob("*.*")))

    def test_scenario_ids_are_unique(self):
        path = OUTPUT / "Tables" / "Scenario_Run_Log.csv"
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        ids = [row["Scenario_ID"] for row in rows]
        self.assertEqual(len(ids), len(set(ids)))

    def test_no_authoritative_input_was_overwritten(self):
        summary = json.loads((OUTPUT / "analysis_summary.json").read_text(encoding="utf-8"))
        self.assertEqual(summary["audit"]["pert_sha256"], summary["authoritative_hashes_after"]["PERT"])
        self.assertEqual(summary["audit"]["gert_sha256"], summary["authoritative_hashes_after"]["GERT"])
