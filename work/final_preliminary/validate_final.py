from __future__ import annotations

import gc
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


OUTPUT_DIR = Path(r"C:\Users\moham\Documents\Codex\2026-07-28\build\outputs")
WORK_DIR = Path(r"C:\Users\moham\Documents\Codex\2026-07-28\build\work\final_preliminary")
WORKBOOK = OUTPUT_DIR / "Final_Preliminary_Simulation_Results.xlsx"
PERT_AUDIT = OUTPUT_DIR / "PERT_Input_Audit.xlsx"
GERT_AUDIT = OUTPUT_DIR / "GERT_Input_Audit.xlsx"
REPORT = OUTPUT_DIR / "Final_Preliminary_Simulation_Report.md"
SUMMARY = OUTPUT_DIR / "Final_Preliminary_Simulation_Summary.csv"
RAW_DIR = OUTPUT_DIR / "Raw_Data"

REQUIRED_SHEETS = [
    "Summary", "Run Info", "Input Hashes", "PERT Input", "PERT Audit", "PERT Summary",
    "PERT Percentiles", "PERT Iterations", "PERT Activities", "PERT Convergence", "PERT Chart Data",
    "PERT Charts", "GERT Nodes", "GERT Arcs", "GERT Audit", "Network Validation", "Exact Verification",
    "GERT Outcomes", "GERT Successful", "GERT Terminated", "GERT Percentiles", "GERT Paths",
    "Node Analysis", "Arc Analysis", "Loop Analysis", "Loop Repetitions", "Cap Events", "Renormalisation",
    "Dispute Analysis", "Transition Analysis", "GERT Iterations", "Reconciliation", "GERT Convergence",
    "Comparison", "GERT Chart Data", "GERT Charts", "Warnings", "Metadata",
]

RAW_FILES = [
    "PERT_Iterations.csv", "PERT_Activity_Durations.csv", "GERT_Iterations.csv", "GERT_Paths.csv",
    "GERT_Node_Visits.csv", "GERT_Arc_Traversals.csv", "GERT_Loops.csv", "GERT_Cap_Events.csv",
    "GERT_Renormalisation.csv", "GERT_Dispute.csv", "GERT_Reconciliation.csv", "Convergence_Checkpoints.csv",
]


def validate() -> dict:
    checks = {}
    checks["exists"] = WORKBOOK.exists()
    checks["larger_than_50kb"] = WORKBOOK.stat().st_size > 50 * 1024
    checks["is_zipfile"] = zipfile.is_zipfile(WORKBOOK)
    if not all(checks.values()): raise AssertionError(checks)

    with zipfile.ZipFile(WORKBOOK) as archive:
        checks["zip_test"] = archive.testzip() is None
        names = archive.namelist()
        chart_files = sorted(name for name in names if name.startswith("xl/charts/chart") and name.endswith(".xml"))
        checks["chart_count"] = len(chart_files) == 26
        checks["no_external_link_parts"] = not any(name.startswith("xl/externalLinks/") for name in names)
        rel_text = "\n".join(archive.read(name).decode("utf-8", errors="ignore") for name in names if name.endswith(".rels"))
        checks["no_external_relationships"] = 'TargetMode="External"' not in rel_text
        chart_texts = [archive.read(name).decode("utf-8", errors="ignore") for name in chart_files]
        checks["every_chart_has_series"] = all(re.search(r"<(?:c:)?ser>", text) for text in chart_texts)
        formulas = [formula for text in chart_texts for formula in re.findall(r"<(?:c:)?f>(.*?)</(?:c:)?f>", text)]
        checks["chart_series_formulas_present"] = len(formulas) >= 26
        checks["chart_ranges_internal"] = all("[" not in formula and "]" not in formula and "!" in formula for formula in formulas)
        xml_text = "\n".join(archive.read(name).decode("utf-8", errors="ignore") for name in names if name.endswith(".xml"))
        checks["no_ref_formulas"] = "#REF!" not in xml_text

    wb = load_workbook(WORKBOOK, read_only=False, data_only=False, keep_links=True)
    checks["required_sheets"] = wb.sheetnames == REQUIRED_SHEETS
    checks["no_empty_required_sheets"] = all(ws.max_row >= 1 and ws.max_column >= 1 for ws in wb.worksheets)
    dimensions = {name: wb[name].calculate_dimension() for name in REQUIRED_SHEETS}
    chart_objects = sum(len(ws._charts) for ws in wb.worksheets)
    checks["openpyxl_chart_objects"] = chart_objects == 26
    checks["openpyxl_external_links"] = len(wb._external_links) == 0
    wb.close()
    del wb
    gc.collect()

    wb2 = load_workbook(WORKBOOK, read_only=False, data_only=False, keep_links=True)
    checks["second_open"] = wb2.sheetnames == REQUIRED_SHEETS and wb2["Summary"].max_row > 1
    wb2.close()
    del wb2
    gc.collect()

    for audit in (PERT_AUDIT, GERT_AUDIT):
        audit_wb = load_workbook(audit, read_only=False, data_only=False)
        if not audit_wb.sheetnames or any(ws.max_row < 1 for ws in audit_wb.worksheets):
            raise AssertionError(f"Invalid audit workbook: {audit}")
        audit_wb.close()
    checks["audit_workbooks_open"] = True
    checks["report_nonempty"] = REPORT.exists() and REPORT.stat().st_size > 0
    checks["summary_nonempty"] = SUMMARY.exists() and SUMMARY.stat().st_size > 0
    checks["raw_files"] = all((RAW_DIR / name).exists() and (RAW_DIR / name).stat().st_size > 0 for name in RAW_FILES)

    failed = [name for name, passed in checks.items() if passed is not True]
    if failed:
        raise AssertionError("Integrity checks failed: " + ", ".join(failed))
    result = {
        "status": "PASS", "checks": checks, "workbook_size_bytes": WORKBOOK.stat().st_size,
        "native_charts": len(chart_files), "chart_series_formulas": len(formulas), "dimensions": dimensions,
        "failed": 0, "errors": 0, "mandatory_skipped": 0,
    }
    (WORK_DIR / "final_integrity.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


if __name__ == "__main__":
    print(json.dumps(validate(), indent=2))
