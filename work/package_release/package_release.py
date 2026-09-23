from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import importlib.util
import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import zipfile
import zlib
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath

import docx
import numpy as np
import openpyxl
import xlsxwriter


BASE = Path(r"C:\Users\moham\Documents\Codex\2026-07-28\build")
OUTPUTS = BASE / "outputs"
SIM = OUTPUTS / "Simulation_Data"
RELEASE = BASE / "release"
PACKAGE_NAME = "VO_PERT_GERT_Full_Simulation_Package"
STAGE = RELEASE / PACKAGE_NAME
ZIP_PATH = RELEASE / f"{PACKAGE_NAME}.zip"
WORK = BASE / "work" / "package_release"
PERT_INPUT = Path(r"C:\Users\moham\OneDrive\Desktop\PERT Input.docx")
GERT_INPUT = Path(r"C:\Users\moham\OneDrive\Desktop\GERT nput.docx")
INTEGRITY_PATH = BASE / "work" / "raw_data_export" / "raw_export_integrity.json"
RAW_MANIFEST = SIM / "RAW_DATA_EXPORT_MANIFEST.csv"
RAW_REPORT = SIM / "RAW_DATA_EXPORT_REPORT.md"
CORE_SOURCE = BASE / "work" / "final_preliminary" / "simulate_final.py"
EXPORT_SOURCE = BASE / "work" / "raw_data_export" / "export_raw_data.py"
PARQUET_SOURCE = BASE / "work" / "raw_data_export" / "ParquetCsvConverter.cs"
PARQUET_HELPER = WORK / "ParquetMetadataCheck.exe"
PARQUET_RUNTIME = BASE / "work" / "raw_data_export" / "parquet_tool"
PROGRAM_VERSION = "raw-data-export-1.0.0"
EXPECTED_PERT_HASH = "7984b5b2e29928345772f0110ea1fa1c9ceb4121dc2e3aed4737112a98bc1f58"
EXPECTED_GERT_HASH = "dc34b00fcd3988146da0760ac18499eeca9d8596601c09445ca5d956c4ae6387"
GENERATED_SELF_REFERENTIAL = {
    "00_START_HERE/PACKAGE_CONTENTS.xlsx",
    "00_START_HERE/PACKAGE_CONTENTS.csv",
    "00_START_HERE/checksums.sha256",
}

FOLDERS = [
    "00_START_HERE",
    "01_AUTHORITATIVE_INPUTS",
    "02_MASTER_AND_METADATA",
    "03_PERT_RESULTS",
    "04_GERT_RESULTS",
    "05_EXACT_VERIFICATION",
    "06_CONVERGENCE",
    "07_SENSITIVITY_READY",
    "08_VALIDATION_AND_RECONCILIATION",
    "09_ALL_CSV",
    "10_ALL_PARQUET",
    "11_SOURCE_AND_REPRODUCIBILITY",
    "12_SUPPORTING_REPORTS",
]

WORKBOOK_DESTINATIONS = {
    "Input_Validation_Report.xlsx": "01_AUTHORITATIVE_INPUTS",
    "00_Master_Index.xlsx": "02_MASTER_AND_METADATA",
    "01_Input_and_Metadata.xlsx": "02_MASTER_AND_METADATA",
    "Data_Dictionary.xlsx": "02_MASTER_AND_METADATA",
    "02_PERT_Raw_Data.xlsx": "03_PERT_RESULTS",
    "03_PERT_Analysis_Ready.xlsx": "03_PERT_RESULTS",
    "04_GERT_Iteration_Data.xlsx": "04_GERT_RESULTS",
    "05_GERT_Traversal_Data.xlsx": "04_GERT_RESULTS",
    "06_GERT_Structure_Events.xlsx": "04_GERT_RESULTS",
    "07_GERT_Analysis_Ready.xlsx": "04_GERT_RESULTS",
    "08_Exact_Verification.xlsx": "05_EXACT_VERIFICATION",
    "09_Convergence_Data.xlsx": "06_CONVERGENCE",
    "10_Sensitivity_Ready_Data.xlsx": "07_SENSITIVITY_READY",
    "11_Validation_and_Reconciliation.xlsx": "08_VALIDATION_AND_RECONCILIATION",
}

DATASET_DESTINATIONS = {
    "00_": "02_MASTER_AND_METADATA",
    "01_": "02_MASTER_AND_METADATA",
    "Data_Dictionary__": "02_MASTER_AND_METADATA",
    "02_": "03_PERT_RESULTS",
    "03_": "03_PERT_RESULTS",
    "04_": "04_GERT_RESULTS",
    "05_": "04_GERT_RESULTS",
    "06_": "04_GERT_RESULTS",
    "07_": "04_GERT_RESULTS",
    "08_": "05_EXACT_VERIFICATION",
    "09_": "06_CONVERGENCE",
    "10_": "07_SENSITIVITY_READY",
    "11_": "08_VALIDATION_AND_RECONCILIATION",
    "Input_Validation_Report__": "08_VALIDATION_AND_RECONCILIATION",
}

DESCRIPTIONS = {
    "00_START_HERE": "Package orientation, inventories, checksums, and validation evidence.",
    "01_AUTHORITATIVE_INPUTS": "Exact authoritative Word inputs and input validation evidence.",
    "02_MASTER_AND_METADATA": "Master index, input metadata, data dictionary, and related datasets.",
    "03_PERT_RESULTS": "PERT raw and analysis-ready workbooks and datasets.",
    "04_GERT_RESULTS": "GERT iteration, traversal, structure-event, and analysis-ready data.",
    "05_EXACT_VERIFICATION": "Finite-state solver outputs, residual diagnostics, and acceptance records.",
    "06_CONVERGENCE": "PERT and GERT convergence checkpoints and stability tests.",
    "07_SENSITIVITY_READY": "Baseline matrices and approved-design templates for future sensitivity analyses.",
    "08_VALIDATION_AND_RECONCILIATION": "Reconciliation, probability, loop, cap, path, reproducibility, and workbook checks.",
    "09_ALL_CSV": "Exact copy of the complete verified CSV directory.",
    "10_ALL_PARQUET": "Exact copy of the complete verified Parquet directory.",
    "11_SOURCE_AND_REPRODUCIBILITY": "Scientific source snapshot and reproducibility metadata.",
    "12_SUPPORTING_REPORTS": "Packaging audit and machine-readable validation evidence.",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalized_relative(path: Path, root: Path = STAGE) -> str:
    return path.relative_to(root).as_posix()


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def copy_exact(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    if source.stat().st_size != destination.stat().st_size or sha256_file(source) != sha256_file(destination):
        raise AssertionError(f"Copy verification failed: {source} -> {destination}")


def safe_recreate_stage() -> None:
    expected = (RELEASE / PACKAGE_NAME).resolve()
    actual = STAGE.resolve()
    if actual != expected or actual.parent != RELEASE.resolve():
        raise RuntimeError(f"Unsafe staging target: {actual}")
    if STAGE.exists():
        shutil.rmtree(STAGE)
    for folder in FOLDERS:
        (STAGE / folder).mkdir(parents=True, exist_ok=True)


def load_core():
    spec = importlib.util.spec_from_file_location("verified_scientific_core", CORE_SOURCE)
    if spec is None or spec.loader is None:
        raise RuntimeError("Unable to load verified scientific core")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def classify_discovered(path: Path) -> tuple[str, str, str]:
    resolved = path.resolve()
    selected_files = {p.resolve() for p in SIM.rglob("*") if p.is_file()}
    selected_files.add((OUTPUTS / "Input_Validation_Report.xlsx").resolve())
    selected_files.update({PERT_INPUT.resolve(), GERT_INPUT.resolve(), CORE_SOURCE.resolve(), EXPORT_SOURCE.resolve(), PARQUET_SOURCE.resolve(), INTEGRITY_PATH.resolve()})
    if resolved in selected_files:
        return "INCLUDE", "", "Verified selected export, authoritative input, source snapshot, or integrity evidence"
    text = str(path)
    if path.name.startswith("~$"):
        return "EXCLUDE", "Temporary Office lock file", "Temporary"
    if "__pycache__" in path.parts or path.suffix == ".pyc":
        return "EXCLUDE", "Python cache", "Cache"
    if "New folder" in path.parts or "Prepared_Simulation_Results" in path.parts:
        return "EXCLUDE", "Earlier or incomplete result package, not the selected verified root", "Obsolete result"
    if path.parent == OUTPUTS:
        return "EXCLUDE", "Superseded by the complete selected Simulation_Data export", "Superseded output"
    return "EXCLUDE", "Not required by the verified package scope", "Unselected"


def discover_files() -> list[dict]:
    candidates = set(p for p in OUTPUTS.rglob("*") if p.is_file())
    candidates.update({PERT_INPUT, GERT_INPUT, CORE_SOURCE, EXPORT_SOURCE, PARQUET_SOURCE, INTEGRITY_PATH})
    records = []
    for path in sorted(candidates, key=lambda p: str(p).lower()):
        inclusion, reason, category = classify_discovered(path)
        exists = path.exists()
        records.append({
            "Selected_Output_Root": str(SIM),
            "Discovered_Path": str(path),
            "Relative_Path": str(path.relative_to(BASE)) if path.is_relative_to(BASE) else path.name,
            "File_Name": path.name,
            "File_Category": category,
            "File_Size_Bytes": path.stat().st_size if exists else 0,
            "SHA256": sha256_file(path) if exists and path.stat().st_size else "",
            "Inclusion_Status": inclusion,
            "Exclusion_Reason": reason,
            "Integrity_Status": "PASS" if exists and path.stat().st_size > 0 else "FAIL",
            "Source_Simulation_Version": PROGRAM_VERSION,
        })
    return records


def compile_parquet_helper() -> None:
    source = WORK / "ParquetMetadataCheck.cs"
    csc = Path(r"C:\Windows\Microsoft.NET\Framework64\v4.0.30319\csc.exe")
    parquet_dll = PARQUET_RUNTIME / "ParquetSharp.dll"
    command = [str(csc), "/nologo", "/target:exe", f"/out:{PARQUET_HELPER}", f"/reference:{parquet_dll}", str(source)]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(f"Parquet metadata helper compilation failed:\n{result.stdout}\n{result.stderr}")
    for dependency in PARQUET_RUNTIME.glob("*.dll"):
        shutil.copy2(dependency, WORK / dependency.name)


def read_parquet_metadata(paths: list[Path]) -> dict[str, tuple[int, int, int]]:
    if not PARQUET_HELPER.exists():
        compile_parquet_helper()
    result = subprocess.run([str(PARQUET_HELPER), *map(str, paths)], cwd=WORK, capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(f"Parquet read failed:\n{result.stdout}\n{result.stderr}")
    metadata = {}
    for line in result.stdout.splitlines():
        path, rows, columns, groups = line.rsplit("\t", 3)
        metadata[str(Path(path).resolve())] = (int(rows), int(columns), int(groups))
    return metadata


def count_csv(path: Path) -> tuple[int, int]:
    opener = gzip.open if path.suffix.lower() == ".gz" else open
    with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader, None)
        if header is None:
            return 0, 0
        rows = sum(1 for _ in reader)
    return rows, len(header)


def workbook_open_test(path: Path) -> dict:
    result = {"path": str(path), "zip_integrity": "FAIL", "openpyxl": "FAIL", "external_links": "FAIL", "ref_errors": "FAIL", "status": "FAIL"}
    if not path.exists() or path.stat().st_size == 0 or not zipfile.is_zipfile(path):
        return result
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            return result
        names = archive.namelist()
        result["zip_integrity"] = "PASS"
        result["external_links"] = "PASS" if not any(name.startswith("xl/externalLinks/") for name in names) else "FAIL"
        ref_found = False
        for name in names:
            if name.endswith(".xml"):
                if b"#REF!" in archive.read(name):
                    ref_found = True
                    break
        result["ref_errors"] = "PASS" if not ref_found else "FAIL"
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=True)
    result["sheet_count"] = len(workbook.sheetnames)
    result["sheet_names"] = ";".join(workbook.sheetnames)
    workbook.close()
    result["openpyxl"] = "PASS"
    result["status"] = "PASS" if all(result[key] == "PASS" for key in ("zip_integrity", "openpyxl", "external_links", "ref_errors")) else "FAIL"
    return result


def add_check(checks: list[dict], category: str, name: str, passed: bool, observed, expected) -> None:
    checks.append({
        "Category": category,
        "Check": name,
        "Observed": observed,
        "Expected": expected,
        "Status": "PASS" if passed else "FAIL",
    })


def prevalidate() -> dict:
    WORK.mkdir(parents=True, exist_ok=True)
    checks: list[dict] = []
    failures: list[str] = []
    pert_hash = sha256_file(PERT_INPUT)
    gert_hash = sha256_file(GERT_INPUT)
    add_check(checks, "Input", "PERT input SHA-256", pert_hash == EXPECTED_PERT_HASH, pert_hash, EXPECTED_PERT_HASH)
    add_check(checks, "Input", "GERT input SHA-256", gert_hash == EXPECTED_GERT_HASH, gert_hash, EXPECTED_GERT_HASH)

    core = load_core()
    pert, nodes, arcs, _, _ = core.parse_inputs()
    input_checks = core.validate_inputs(pert, nodes, arcs)
    add_check(checks, "Input", "PERT route activity count", len(pert) == 6, len(pert), 6)
    add_check(checks, "Input", "GERT node count", len(nodes) == 10, len(nodes), 10)
    internal_arc_count = len(arcs) + 1  # The deterministic S0 -> S1 structural arc is stored separately by the core parser.
    add_check(checks, "Input", "GERT internal arc count", internal_arc_count == 30, internal_arc_count, 30)
    e12 = next((arc for arc in arcs if arc["arc_tag"] == "e12"), None)
    add_check(checks, "Input", "e12 mapping", bool(e12 and e12["from_node"] == "S1" and e12["to_node"] == "S2"), f"{e12['from_node']} -> {e12['to_node']}" if e12 else "missing", "S1 -> S2")
    add_check(checks, "Input", "Scientific-core input validation", all(row["status"] == "PASS" for row in input_checks), f"{sum(row['status'] == 'PASS' for row in input_checks)}/{len(input_checks)} PASS", "all PASS")

    integrity = json.loads(INTEGRITY_PATH.read_text(encoding="utf-8"))
    expected_counts = {
        "pert_iterations": 50000,
        "pert_long_rows": 300000,
        "gert_iterations": 50000,
        "gert_traversals": 226832,
        "gert_node_visits": 276832,
        "successful": 28871,
        "terminated": 21129,
        "workbooks_passed": 14,
        "workbooks_failed": 0,
        "tests_failed": 0,
    }
    for key, expected in expected_counts.items():
        observed = integrity["counts"].get(key)
        add_check(checks, "Simulation", key, observed == expected, observed, expected)
    add_check(checks, "Simulation", "Successful + Terminated", integrity["counts"]["successful"] + integrity["counts"]["terminated"] == 50000, integrity["counts"]["successful"] + integrity["counts"]["terminated"], 50000)
    add_check(checks, "Simulation", "Exact verification", integrity["exact_status"] == "PASS", integrity["exact_status"], "PASS")
    add_check(checks, "Simulation", "Reproducibility", integrity["reproducibility"] == "PASS", integrity["reproducibility"], "PASS")
    add_check(checks, "Simulation", "Cross-file reconciliation", integrity["cross_file_reconciliation"] == "PASS", integrity["cross_file_reconciliation"], "PASS")

    manifest_rows = list(csv.DictReader(RAW_MANIFEST.open(encoding="utf-8-sig", newline="")))
    csv_rows = [row for row in manifest_rows if row["Artifact_Type"] == "CSV Dataset"]
    parquet_rows = [row for row in manifest_rows if row["Artifact_Type"] == "Parquet Dataset"]
    workbook_rows = [row for row in manifest_rows if row["Artifact_Type"] == "Excel Workbook"]
    add_check(checks, "Export", "CSV dataset manifest count", len(csv_rows) == 127, len(csv_rows), 127)
    add_check(checks, "Export", "Parquet dataset manifest count", len(parquet_rows) == 127, len(parquet_rows), 127)
    add_check(checks, "Export", "Excel workbook manifest count", len(workbook_rows) == 14, len(workbook_rows), 14)

    csv_validation = []
    for row in csv_rows:
        path = SIM / "CSV" / row["File_Name"]
        actual_rows, actual_columns = count_csv(path)
        status = actual_rows == int(row["Rows"]) and actual_columns == int(row["Columns"]) and path.stat().st_size > 0
        csv_validation.append({"File": row["File_Name"], "Rows": actual_rows, "Columns": actual_columns, "Expected_Rows": int(row["Rows"]), "Expected_Columns": int(row["Columns"]), "Status": "PASS" if status else "FAIL"})
    add_check(checks, "Export", "All CSV files read with expected dimensions", all(row["Status"] == "PASS" for row in csv_validation), sum(row["Status"] == "PASS" for row in csv_validation), 127)

    parquet_paths = [SIM / "Parquet" / row["File_Name"] for row in parquet_rows]
    parquet_metadata = read_parquet_metadata(parquet_paths)
    parquet_validation = []
    for row, path in zip(parquet_rows, parquet_paths):
        actual_rows, actual_columns, row_groups = parquet_metadata[str(path.resolve())]
        status = actual_rows == int(row["Rows"]) and actual_columns == int(row["Columns"]) and row_groups > 0 and path.stat().st_size > 0
        parquet_validation.append({"File": row["File_Name"], "Rows": actual_rows, "Columns": actual_columns, "Row_Groups": row_groups, "Expected_Rows": int(row["Rows"]), "Expected_Columns": int(row["Columns"]), "Status": "PASS" if status else "FAIL"})
    add_check(checks, "Export", "All Parquet files read with expected dimensions", all(row["Status"] == "PASS" for row in parquet_validation), sum(row["Status"] == "PASS" for row in parquet_validation), 127)
    add_check(checks, "Export", "CSV-Parquet parity", all(c["Rows"] == p["Rows"] and c["Columns"] == p["Columns"] for c, p in zip(csv_validation, parquet_validation)), "127 paired datasets", "all dimensions equal")

    workbook_paths = [Path(row["Relative_Path"].replace("outputs\\", str(OUTPUTS) + os.sep, 1)) for row in workbook_rows]
    workbook_validation = [workbook_open_test(path) for path in workbook_paths]
    add_check(checks, "Export", "All workbooks pass ZIP/open/external-link/#REF tests", all(row["status"] == "PASS" for row in workbook_validation), sum(row["status"] == "PASS" for row in workbook_validation), 14)

    all_selected = [PERT_INPUT, GERT_INPUT, RAW_MANIFEST, RAW_REPORT, INTEGRITY_PATH, CORE_SOURCE, EXPORT_SOURCE, PARQUET_SOURCE]
    all_selected.extend(SIM.rglob("*"))
    zero_files = [str(path) for path in all_selected if path.is_file() and path.stat().st_size == 0]
    add_check(checks, "Export", "No selected file has zero bytes", not zero_files, len(zero_files), 0)

    failures = [row["Check"] for row in checks if row["Status"] != "PASS"]
    payload = {
        "status": "PASS" if not failures else "FAIL",
        "timestamp_utc": utc_now(),
        "selected_output_root": str(SIM),
        "simulation_version": PROGRAM_VERSION,
        "checks": checks,
        "failures": failures,
        "discovered_files": discover_files(),
        "csv_validation": csv_validation,
        "parquet_validation": parquet_validation,
        "workbook_validation": workbook_validation,
        "raw_manifest_rows": manifest_rows,
        "integrity": integrity,
        "input_hashes": {"PERT Input.docx": pert_hash, "GERT nput.docx": gert_hash},
    }
    write_json(WORK / "prepackaging_audit.json", payload)
    if failures:
        raise AssertionError(f"Pre-packaging validation failed: {failures}")
    return payload


def dataset_folder(name: str) -> str:
    for prefix, folder in DATASET_DESTINATIONS.items():
        if name.startswith(prefix):
            return folder
    raise KeyError(f"Unmapped dataset file: {name}")


def write_readme() -> None:
    content = """VO PERT-GERT FULL SIMULATION PACKAGE
====================================

This package contains the complete preliminary PERT-MCS and GERT-MCS simulation dataset.

Simulation configuration
------------------------
- 50,000 iterations per model.
- Root seed 42.
- Beta-PERT lambda 4.
- Activity- and traversal-level ceiling.

Recommended starting files
--------------------------
- 02_MASTER_AND_METADATA/00_Master_Index.xlsx
- 00_START_HERE/RAW_DATA_EXPORT_REPORT.md
- 02_MASTER_AND_METADATA/Data_Dictionary.xlsx
- 03_PERT_RESULTS/03_PERT_Analysis_Ready.xlsx
- 04_GERT_RESULTS/07_GERT_Analysis_Ready.xlsx
- 07_SENSITIVITY_READY/10_Sensitivity_Ready_Data.xlsx
- 08_VALIDATION_AND_RECONCILIATION/11_Validation_and_Reconciliation.xlsx

Format guidance
---------------
Excel files are provided for review and analysis. Parquet files are the preferred format for large programmatic analyses. CSV files provide universal fallback access.

Analytical scope
----------------
The baseline matrices support descriptive, correlation, regression, path, structural, and tail-driver analyses. Sensitivity to fixed probabilities, O/ML/P parameters, and Loop Caps requires additional approved parameter-variation runs. Baseline associations must not be interpreted as causal effects.

Integrity
---------
See checksums.sha256, PACKAGE_CONTENTS.xlsx, and PACKAGE_VALIDATION_REPORT.md before analysis. Package-relative paths are used throughout.
"""
    (STAGE / "00_START_HERE" / "README_FIRST.txt").write_text(content, encoding="utf-8", newline="\n")


def write_input_hashes(audit: dict) -> None:
    lines = [f"{digest}  {name}" for name, digest in audit["input_hashes"].items()]
    (STAGE / "01_AUTHORITATIVE_INPUTS" / "input_hashes.txt").write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


def dependency_payload() -> dict:
    return {
        "generated_utc": utc_now(),
        "program_version": PROGRAM_VERSION,
        "runtime": {
            "python": platform.python_version(),
            "implementation": platform.python_implementation(),
            "platform": platform.platform(),
        },
        "python_packages": {
            "numpy": np.__version__,
            "openpyxl": openpyxl.__version__,
            "XlsxWriter": xlsxwriter.__version__,
            "python-docx": docx.__version__,
        },
        "parquet_writer": {
            "implementation": "ParquetSharp",
            "compression": "Snappy",
            "row_group_size": 10000,
            "runtime_files": [
                {"name": path.name, "sha256": sha256_file(path), "size_bytes": path.stat().st_size}
                for path in sorted(PARQUET_RUNTIME.glob("*.dll"))
            ],
        },
        "source_files": [
            {"name": source.name, "sha256": sha256_file(source), "size_bytes": source.stat().st_size}
            for source in (CORE_SOURCE, EXPORT_SOURCE, PARQUET_SOURCE)
        ],
    }


def write_reproducibility_materials(audit: dict) -> None:
    target = STAGE / "11_SOURCE_AND_REPRODUCIBILITY"
    source_root = target / "Source"
    copy_exact(CORE_SOURCE, source_root / "simulation_core" / CORE_SOURCE.name)
    copy_exact(EXPORT_SOURCE, source_root / "export_generation" / EXPORT_SOURCE.name)
    copy_exact(PARQUET_SOURCE, source_root / "parquet_export" / PARQUET_SOURCE.name)
    dependency = dependency_payload()
    write_json(target / "Dependency_Manifest.json", dependency)
    requirements = [
        f"numpy=={np.__version__}",
        f"openpyxl=={openpyxl.__version__}",
        f"python-docx=={docx.__version__}",
        f"XlsxWriter=={xlsxwriter.__version__}",
    ]
    (target / "requirements-lock.txt").write_text("\n".join(requirements) + "\n", encoding="utf-8", newline="\n")
    (target / "PROGRAM_VERSION.txt").write_text(PROGRAM_VERSION + "\n", encoding="utf-8", newline="\n")
    (target / "GIT_COMMIT.txt").write_text("Not available: the selected workspace is not a Git repository.\n", encoding="utf-8", newline="\n")
    write_json(target / "simulation_configuration.json", {
        "model": "VO PERT-MCS and GERT-MCS",
        "pert_iterations": 50000,
        "gert_iterations": 50000,
        "root_seed": 42,
        "beta_pert_lambda": 4.0,
        "rounding_rule": "ceil at activity and traversal level",
        "pert_input_sha256": audit["input_hashes"]["PERT Input.docx"],
        "gert_input_sha256": audit["input_hashes"]["GERT nput.docx"],
        "program_version": PROGRAM_VERSION,
    })
    write_json(target / "random_stream_metadata.json", {
        "generator": "NumPy default_rng (PCG64)",
        "root_seed": 42,
        "seed_sequence_spawn_count": 3,
        "streams": [
            {"spawn_index": 0, "purpose": "PERT beta sampling"},
            {"spawn_index": 1, "purpose": "GERT routing"},
            {"spawn_index": 2, "purpose": "GERT duration sampling"},
        ],
        "reproducibility_status": audit["integrity"]["reproducibility"],
    })


def stage(audit: dict) -> None:
    safe_recreate_stage()
    write_readme()
    copy_exact(RAW_REPORT, STAGE / "00_START_HERE" / RAW_REPORT.name)
    copy_exact(RAW_MANIFEST, STAGE / "00_START_HERE" / RAW_MANIFEST.name)
    copy_exact(PERT_INPUT, STAGE / "01_AUTHORITATIVE_INPUTS" / PERT_INPUT.name)
    copy_exact(GERT_INPUT, STAGE / "01_AUTHORITATIVE_INPUTS" / GERT_INPUT.name)
    copy_exact(OUTPUTS / "Input_Validation_Report.xlsx", STAGE / "01_AUTHORITATIVE_INPUTS" / "Input_Validation_Report.xlsx")
    write_input_hashes(audit)

    for name, folder in WORKBOOK_DESTINATIONS.items():
        if name == "Input_Validation_Report.xlsx":
            continue
        copy_exact(SIM / name, STAGE / folder / name)

    for source in sorted((SIM / "CSV").iterdir()):
        if source.is_file():
            copy_exact(source, STAGE / dataset_folder(source.name) / "CSV" / source.name)
            copy_exact(source, STAGE / "09_ALL_CSV" / source.name)
    for source in sorted((SIM / "Parquet").iterdir()):
        if source.is_file():
            copy_exact(source, STAGE / dataset_folder(source.name) / "Parquet" / source.name)
            copy_exact(source, STAGE / "10_ALL_PARQUET" / source.name)

    write_reproducibility_materials(audit)
    copy_exact(OUTPUTS / "Packaging_Audit.xlsx", STAGE / "12_SUPPORTING_REPORTS" / "Packaging_Audit.xlsx")
    copy_exact(INTEGRITY_PATH, STAGE / "12_SUPPORTING_REPORTS" / "raw_export_integrity.json")
    copy_exact(WORK / "prepackaging_audit.json", STAGE / "12_SUPPORTING_REPORTS" / "prepackaging_validation.json")

    validation_placeholder = "# Package Validation Report\n\nStatus: PRELIMINARY\n\nThe completed report is inserted after preliminary ZIP extraction and integrity testing.\n"
    (STAGE / "00_START_HERE" / "PACKAGE_VALIDATION_REPORT.md").write_text(validation_placeholder, encoding="utf-8", newline="\n")
    build_manifest_seed(audit, phase="preliminary")


def description_for(path: Path) -> str:
    relative = normalized_relative(path)
    top = PurePosixPath(relative).parts[0]
    if path.suffix.lower() == ".xlsx":
        return "Excel workbook for review, analysis, audit, or package inventory."
    if path.suffix.lower() == ".parquet":
        return "Verified Parquet dataset for programmatic analysis."
    if path.name.endswith(".csv") or path.name.endswith(".csv.gz"):
        return "Verified CSV dataset or report for universal access."
    if path.suffix.lower() == ".docx":
        return "Authoritative model input document."
    if top == "11_SOURCE_AND_REPRODUCIBILITY":
        return "Scientific source or reproducibility metadata."
    return DESCRIPTIONS.get(top, "Packaged simulation support file.")


def category_for(relative: str) -> str:
    return DESCRIPTIONS.get(PurePosixPath(relative).parts[0], "Package file")


def original_source_for(relative: str) -> str:
    name = PurePosixPath(relative).name
    if relative.startswith("01_AUTHORITATIVE_INPUTS/"):
        if name == PERT_INPUT.name:
            return str(PERT_INPUT)
        if name == GERT_INPUT.name:
            return str(GERT_INPUT)
        if name == "Input_Validation_Report.xlsx":
            return str(OUTPUTS / name)
    if name in WORKBOOK_DESTINATIONS and name != "Input_Validation_Report.xlsx":
        return str(SIM / name)
    if "/CSV/" in relative or relative.startswith("09_ALL_CSV/"):
        return str(SIM / "CSV" / name)
    if "/Parquet/" in relative or relative.startswith("10_ALL_PARQUET/"):
        return str(SIM / "Parquet" / name)
    if name == RAW_REPORT.name:
        return str(RAW_REPORT)
    if name == RAW_MANIFEST.name:
        return str(RAW_MANIFEST)
    if name == CORE_SOURCE.name:
        return str(CORE_SOURCE)
    if name == EXPORT_SOURCE.name:
        return str(EXPORT_SOURCE)
    if name == PARQUET_SOURCE.name:
        return str(PARQUET_SOURCE)
    if name == "Packaging_Audit.xlsx":
        return str(OUTPUTS / name)
    if name == "raw_export_integrity.json":
        return str(INTEGRITY_PATH)
    return "Generated during release packaging"


def dataset_lookup(audit: dict) -> dict[tuple[str, str], dict]:
    lookup = {}
    for row in audit["raw_manifest_rows"]:
        if row["Artifact_Type"] in {"CSV Dataset", "Parquet Dataset"}:
            lookup[(row["Artifact_Type"], row["File_Name"])] = row
    return lookup


def build_manifest_seed(audit: dict, phase: str) -> None:
    lookup = dataset_lookup(audit)
    expected_generated = [
        STAGE / "00_START_HERE" / "PACKAGE_CONTENTS.xlsx",
        STAGE / "00_START_HERE" / "PACKAGE_CONTENTS.csv",
        STAGE / "00_START_HERE" / "checksums.sha256",
    ]
    paths = [path for path in STAGE.rglob("*") if path.is_file()]
    for path in expected_generated:
        if path not in paths:
            paths.append(path)
    records = []
    for path in sorted(paths, key=lambda p: normalized_relative(p).lower()):
        relative = normalized_relative(path)
        exists = path.exists()
        is_csv_dataset = ("/CSV/" in relative or relative.startswith("09_ALL_CSV/")) and (path.name.endswith(".csv") or path.name.endswith(".csv.gz"))
        is_parquet_dataset = ("/Parquet/" in relative or relative.startswith("10_ALL_PARQUET/")) and path.suffix.lower() == ".parquet"
        artifact_type = "CSV Dataset" if is_csv_dataset else "Parquet Dataset" if is_parquet_dataset else "File"
        source_manifest = lookup.get((artifact_type, path.name), {})
        self_reference = relative in GENERATED_SELF_REFERENTIAL
        records.append({
            "Package_Relative_Path": relative,
            "File_Name": path.name,
            "Category": category_for(relative),
            "Description": description_for(path),
            "File_Size_Bytes": path.stat().st_size if exists else None,
            "SHA256": "SELF-REFERENTIAL; see checksums.sha256" if self_reference else sha256_file(path) if exists else "PENDING",
            "Original_Source_Path": original_source_for(relative),
            "Required_or_Optional": "Required",
            "Integrity_Status": "PASS" if exists and path.stat().st_size > 0 else "PENDING",
            "Open_Test_Status": "PASS" if exists and path.suffix.lower() == ".xlsx" and path.name != "PACKAGE_CONTENTS.xlsx" else "N/A" if exists else "PENDING",
            "Dataset": source_manifest.get("Dataset", ""),
            "Dataset_Row_Count": int(source_manifest["Rows"]) if source_manifest.get("Rows") else None,
            "Dataset_Column_Count": int(source_manifest["Columns"]) if source_manifest.get("Columns") else None,
            "Unique_Dataset_Key": source_manifest.get("Dataset", ""),
            "Manifest_Phase": phase,
        })
    source_files = [row for row in records if row["Package_Relative_Path"].startswith("11_SOURCE_AND_REPRODUCIBILITY/Source/")]
    dataset_records = [row for row in records if row["Dataset"]]
    summary = {
        "package_name": PACKAGE_NAME,
        "phase": phase,
        "generated_utc": utc_now(),
        "selected_output_root": str(SIM),
        "total_packaged_files": len(records),
        "total_folders": len({str(PurePosixPath(row["Package_Relative_Path"]).parent) for row in records}),
        "excel_workbooks": sum(row["File_Name"].lower().endswith(".xlsx") for row in records),
        "csv_dataset_file_copies": sum(row["Package_Relative_Path"].endswith(".csv") or row["Package_Relative_Path"].endswith(".csv.gz") for row in dataset_records),
        "parquet_dataset_file_copies": sum(row["Package_Relative_Path"].endswith(".parquet") for row in dataset_records),
        "unique_csv_datasets": len({row["Unique_Dataset_Key"] for row in dataset_records if row["Package_Relative_Path"].endswith((".csv", ".csv.gz"))}),
        "unique_parquet_datasets": len({row["Unique_Dataset_Key"] for row in dataset_records if row["Package_Relative_Path"].endswith(".parquet")}),
        "source_files": len(source_files),
        "total_uncompressed_size": sum(row["File_Size_Bytes"] or 0 for row in records),
        "validation_status": audit["status"],
        "missing_files": sum(row["Integrity_Status"] == "PENDING" for row in records),
    }
    payload = {"summary": summary, "files": records, "datasets": dataset_records, "prevalidation_checks": audit["checks"]}
    write_json(WORK / "package_manifest_seed.json", payload)


def create_checksums() -> None:
    checksum_path = STAGE / "00_START_HERE" / "checksums.sha256"
    paths = [path for path in STAGE.rglob("*") if path.is_file() and path != checksum_path]
    lines = [f"{sha256_file(path)}  {normalized_relative(path)}" for path in sorted(paths, key=lambda p: normalized_relative(p).lower())]
    checksum_path.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


def zip_stage() -> None:
    RELEASE.mkdir(parents=True, exist_ok=True)
    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6, allowZip64=True) as archive:
        for path in sorted((p for p in STAGE.rglob("*") if p.is_file()), key=lambda p: normalized_relative(p).lower()):
            arcname = f"{PACKAGE_NAME}/{normalized_relative(path)}"
            archive.write(path, arcname=arcname)


def parse_checksums(path: Path) -> dict[str, str]:
    result = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        digest, relative = line.split("  ", 1)
        result[relative] = digest
    return result


def zip_member_crc(path: Path) -> int:
    crc = 0
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            crc = zlib.crc32(block, crc)
    return crc & 0xFFFFFFFF


def validate_zip(label: str, keep_details: bool = True) -> dict:
    checks = []
    corrupt_members = []
    missing_members = []
    size_mismatches = []
    crc_mismatches = []
    hash_mismatches = []
    excel_failures = []
    csv_failures = []
    parquet_failures = []
    staged_files = [p for p in STAGE.rglob("*") if p.is_file()]
    expected_members = {f"{PACKAGE_NAME}/{normalized_relative(path)}": path for path in staged_files}

    is_zip = ZIP_PATH.exists() and ZIP_PATH.stat().st_size > 0 and zipfile.is_zipfile(ZIP_PATH)
    add_check(checks, "ZIP", "ZIP exists, nonzero, and is valid format", is_zip, ZIP_PATH.stat().st_size if ZIP_PATH.exists() else 0, ">0 and zipfile=True")
    if not is_zip:
        raise AssertionError("ZIP file is missing or invalid")
    with zipfile.ZipFile(ZIP_PATH) as archive:
        corrupt = archive.testzip()
        if corrupt:
            corrupt_members.append(corrupt)
        infos = {info.filename: info for info in archive.infolist() if not info.is_dir()}
        missing_members = sorted(set(expected_members) - set(infos))
        unexpected_members = sorted(set(infos) - set(expected_members))
        for member, path in expected_members.items():
            if member not in infos:
                continue
            info = infos[member]
            if info.file_size != path.stat().st_size:
                size_mismatches.append(member)
            if info.CRC != zip_member_crc(path):
                crc_mismatches.append(member)
    add_check(checks, "ZIP", "ZipFile.testzip", not corrupt_members, len(corrupt_members), 0)
    add_check(checks, "ZIP", "Staged/ZIP member path parity", not missing_members and not unexpected_members, f"missing={len(missing_members)} unexpected={len(unexpected_members)}", "0/0")
    add_check(checks, "ZIP", "Uncompressed size parity", not size_mismatches, len(size_mismatches), 0)
    add_check(checks, "ZIP", "CRC parity", not crc_mismatches, len(crc_mismatches), 0)

    verification_root = Path(tempfile.mkdtemp(prefix="vo_pert_gert_zip_verify_", dir=WORK))
    try:
        with zipfile.ZipFile(ZIP_PATH) as archive:
            archive.extractall(verification_root)
        extracted = verification_root / PACKAGE_NAME
        checksum_map = parse_checksums(extracted / "00_START_HERE" / "checksums.sha256")
        for relative, expected_hash in checksum_map.items():
            path = extracted / Path(relative)
            if not path.exists() or sha256_file(path) != expected_hash:
                hash_mismatches.append(relative)
        add_check(checks, "Extraction", "Extracted hashes match checksums manifest", not hash_mismatches, len(hash_mismatches), 0)

        extracted_files = [p for p in extracted.rglob("*") if p.is_file()]
        required_names = set(WORKBOOK_DESTINATIONS) | {
            "PERT Input.docx", "GERT nput.docx", "PACKAGE_CONTENTS.xlsx", "README_FIRST.txt",
            "RAW_DATA_EXPORT_REPORT.md", "RAW_DATA_EXPORT_MANIFEST.csv", "PACKAGE_VALIDATION_REPORT.md",
        }
        extracted_names = {p.name for p in extracted_files}
        missing_required = sorted(required_names - extracted_names)
        add_check(checks, "Extraction", "Required file names present", not missing_required, len(missing_required), 0)

        for path in sorted(p for p in extracted_files if p.suffix.lower() == ".xlsx"):
            try:
                result = workbook_open_test(path)
                if result["status"] != "PASS":
                    excel_failures.append(normalized_relative(path, extracted))
            except Exception as exc:
                excel_failures.append(f"{normalized_relative(path, extracted)}: {exc}")
        add_check(checks, "Extraction", "All extracted Excel workbooks open", not excel_failures, len(excel_failures), 0)

        csv_sample = extracted / "09_ALL_CSV" / "00_Master_Index__Simulation_Summary.csv"
        try:
            rows, columns = count_csv(csv_sample)
            if rows < 1 or columns < 1:
                csv_failures.append(str(csv_sample))
        except Exception as exc:
            csv_failures.append(str(exc))
        add_check(checks, "Extraction", "Extracted CSV sample read", not csv_failures, len(csv_failures), 0)

        parquet_sample = extracted / "10_ALL_PARQUET" / "00_Master_Index__Simulation_Summary.parquet"
        try:
            metadata = read_parquet_metadata([parquet_sample])
            rows, columns, groups = metadata[str(parquet_sample.resolve())]
            if rows < 1 or columns < 1 or groups < 1:
                parquet_failures.append(str(parquet_sample))
        except Exception as exc:
            parquet_failures.append(str(exc))
        add_check(checks, "Extraction", "Extracted Parquet sample read", not parquet_failures, len(parquet_failures), 0)

        csv_count = sum(1 for p in extracted_files if (p.name.endswith(".csv") or p.name.endswith(".csv.gz")) and ("CSV" in p.parts or "09_ALL_CSV" in p.parts))
        parquet_count = sum(1 for p in extracted_files if p.suffix.lower() == ".parquet")
        excel_count = sum(1 for p in extracted_files if p.suffix.lower() == ".xlsx")
        source_count = sum(1 for p in extracted_files if "Source" in p.parts)
        folder_count = len({p.parent for p in extracted_files})
        add_check(checks, "Extraction", "Physical CSV dataset copies", csv_count == 254, csv_count, 254)
        add_check(checks, "Extraction", "Physical Parquet dataset copies", parquet_count == 254, parquet_count, 254)
        add_check(checks, "Extraction", "Unique CSV datasets", len({p.name for p in extracted_files if p.name.endswith((".csv", ".csv.gz")) and ("CSV" in p.parts or "09_ALL_CSV" in p.parts)}) == 127, len({p.name for p in extracted_files if p.name.endswith((".csv", ".csv.gz")) and ("CSV" in p.parts or "09_ALL_CSV" in p.parts)}), 127)
        add_check(checks, "Extraction", "Unique Parquet datasets", len({p.name for p in extracted_files if p.suffix.lower() == ".parquet"}) == 127, len({p.name for p in extracted_files if p.suffix.lower() == ".parquet"}), 127)

        payload = {
            "label": label,
            "status": "PASS" if all(row["Status"] == "PASS" for row in checks) else "FAIL",
            "validated_utc": utc_now(),
            "zip_path": str(ZIP_PATH),
            "zip_size_bytes": ZIP_PATH.stat().st_size,
            "zip_sha256": sha256_file(ZIP_PATH),
            "total_packaged_files": len(extracted_files),
            "total_package_folders": folder_count,
            "excel_workbooks": excel_count,
            "unique_csv_datasets": 127,
            "unique_parquet_datasets": 127,
            "csv_dataset_file_copies": csv_count,
            "parquet_dataset_file_copies": parquet_count,
            "total_raw_data_file_copies": csv_count + parquet_count,
            "source_files": source_count,
            "total_uncompressed_size": sum(p.stat().st_size for p in extracted_files),
            "corrupt_zip_members": len(corrupt_members),
            "missing_required_files": len(missing_required),
            "hash_mismatches_after_extraction": len(hash_mismatches),
            "excel_open_failures_after_extraction": len(excel_failures),
            "csv_read_failures": len(csv_failures),
            "parquet_read_failures": len(parquet_failures),
            "checks": checks,
            "details": {
                "corrupt_members": corrupt_members,
                "missing_members": missing_members,
                "size_mismatches": size_mismatches,
                "crc_mismatches": crc_mismatches,
                "hash_mismatches": hash_mismatches,
                "excel_failures": excel_failures,
                "csv_failures": csv_failures,
                "parquet_failures": parquet_failures,
            },
        }
        if payload["status"] != "PASS":
            raise AssertionError(f"ZIP validation failed: {payload['details']}")
        if keep_details:
            write_json(WORK / f"{label}_zip_validation.json", payload)
        return payload
    finally:
        expected_parent = WORK.resolve()
        resolved = verification_root.resolve()
        if resolved.parent == expected_parent and resolved.name.startswith("vo_pert_gert_zip_verify_"):
            shutil.rmtree(resolved)


def write_validation_report(preliminary: dict, final: dict | None = None) -> None:
    source = final or preliminary
    final_note = "Final validation completed after the report was inserted and the ZIP was recreated." if final else "The completed full extraction and integrity validation below was performed on the pre-report archive. This report was then inserted, the master archive was recreated, and the recreated archive was independently revalidated before delivery."
    lines = [
        "# Package Validation Report",
        "",
        f"Status: {source['status']}",
        f"Validation timestamp (UTC): {source['validated_utc']}",
        "",
        final_note,
        "",
        "## Package Counts",
        "",
        f"- Total packaged files: {source['total_packaged_files']}",
        f"- Total package folders: {source['total_package_folders']}",
        f"- Total Excel workbooks: {source['excel_workbooks']}",
        f"- Unique CSV datasets: {source['unique_csv_datasets']}",
        f"- Unique Parquet datasets: {source['unique_parquet_datasets']}",
        f"- Physical CSV dataset copies: {source['csv_dataset_file_copies']}",
        f"- Physical Parquet dataset copies: {source['parquet_dataset_file_copies']}",
        f"- Total raw-data file copies: {source['total_raw_data_file_copies']}",
        f"- Total source files: {source['source_files']}",
        f"- Total uncompressed size: {source['total_uncompressed_size']} bytes",
        f"- Validated pre-report ZIP size: {source['zip_size_bytes']} bytes",
        f"- Validated pre-report ZIP SHA-256: {source['zip_sha256']}",
        "",
        "The organized result folders contain one copy of each of the 127 CSV and 127 Parquet datasets. The complete verified CSV and Parquet directories are also copied unchanged into 09_ALL_CSV and 10_ALL_PARQUET, producing 254 physical files per format without changing dataset counts.",
        "",
        "The recreated master ZIP size and SHA-256 are reported with release delivery. A ZIP cannot contain its own stable final SHA-256 because inserting that digest would change the archive bytes.",
        "",
        "## Integrity Results",
        "",
        f"- Corrupt ZIP members: {source['corrupt_zip_members']}",
        f"- Missing required files: {source['missing_required_files']}",
        f"- Hash mismatches after extraction: {source['hash_mismatches_after_extraction']}",
        f"- Excel open failures after extraction: {source['excel_open_failures_after_extraction']}",
        f"- CSV read failures: {source['csv_read_failures']}",
        f"- Parquet read failures: {source['parquet_read_failures']}",
        "",
        "## Validation Checks",
        "",
        "| Category | Check | Observed | Expected | Status |",
        "|---|---|---:|---:|---|",
    ]
    for check in source["checks"]:
        observed = str(check["Observed"]).replace("|", "\\|")
        expected = str(check["Expected"]).replace("|", "\\|")
        lines.append(f"| {check['Category']} | {check['Check']} | {observed} | {expected} | {check['Status']} |")
    lines.extend([
        "",
        "## Manifest Note",
        "",
        "checksums.sha256 is the detached machine-verification manifest for every packaged file except itself. PACKAGE_CONTENTS.xlsx and PACKAGE_CONTENTS.csv record every packaged path; their own SHA-256 cells are marked self-referential because embedding a file's final cryptographic digest inside that same file is not mathematically stable.",
        "",
    ])
    (STAGE / "00_START_HERE" / "PACKAGE_VALIDATION_REPORT.md").write_text("\n".join(lines), encoding="utf-8", newline="\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["audit", "stage", "manifest", "zip", "validate-preliminary", "write-report", "validate-final"])
    args = parser.parse_args()
    if args.command == "audit":
        payload = prevalidate()
        print(json.dumps({"status": payload["status"], "checks": len(payload["checks"]), "discovered": len(payload["discovered_files"])}))
    elif args.command == "stage":
        audit = json.loads((WORK / "prepackaging_audit.json").read_text(encoding="utf-8"))
        stage(audit)
        print(json.dumps({"staged_files": sum(1 for p in STAGE.rglob('*') if p.is_file())}))
    elif args.command == "manifest":
        audit = json.loads((WORK / "prepackaging_audit.json").read_text(encoding="utf-8"))
        phase = "final" if (WORK / "preliminary_zip_validation.json").exists() else "preliminary"
        build_manifest_seed(audit, phase)
        print(json.dumps(json.loads((WORK / "package_manifest_seed.json").read_text(encoding="utf-8"))["summary"]))
    elif args.command == "zip":
        create_checksums()
        zip_stage()
        print(json.dumps({"zip": str(ZIP_PATH), "size": ZIP_PATH.stat().st_size, "sha256": sha256_file(ZIP_PATH)}))
    elif args.command == "validate-preliminary":
        payload = validate_zip("preliminary")
        print(json.dumps(payload))
    elif args.command == "write-report":
        preliminary = json.loads((WORK / "preliminary_zip_validation.json").read_text(encoding="utf-8"))
        write_validation_report(preliminary)
        print(str(STAGE / "00_START_HERE" / "PACKAGE_VALIDATION_REPORT.md"))
    elif args.command == "validate-final":
        payload = validate_zip("final")
        write_json(WORK / "final_package_summary.json", payload)
        print(json.dumps(payload))


if __name__ == "__main__":
    main()
