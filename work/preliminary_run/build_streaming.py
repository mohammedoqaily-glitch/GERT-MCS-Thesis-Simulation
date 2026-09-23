from __future__ import annotations

import json
import math
import os
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


WORK_DIR = Path(r"C:\Users\moham\Documents\Codex\2026-07-28\build\work\preliminary_run")
OUTPUT_DIR = Path(r"C:\Users\moham\Documents\Codex\2026-07-28\build\outputs")
DATA = json.loads((WORK_DIR / "simulation_data.json").read_text(encoding="utf-8"))

MAIN_PATH = OUTPUT_DIR / "Preliminary_Experimental_Simulation_Results.xlsx"
REPORT_PATH = OUTPUT_DIR / "Preliminary_Experimental_Simulation_Report.md"
AUDIT_PATH = OUTPUT_DIR / "Preliminary_Input_Audit.xlsx"

REQUIRED_SHEETS = [
    "Executive Summary", "Run Information", "Input Audit", "Word Parsed Records", "Validation",
    "PERT Input", "PERT Summary", "PERT Analytical Check", "PERT P1-P99", "PERT Iterations",
    "PERT Activity Durations", "PERT Convergence", "PERT Histogram Data", "PERT ECDF Data",
    "PERT Chart Data", "PERT Charts", "GERT Status", "GERT Input Audit", "GERT Outcome Summary",
    "GERT Successful Results", "GERT Terminated Results", "GERT Paths", "GERT Nodes", "GERT Arcs",
    "GERT Loops", "GERT Cap Events", "GERT Renormalisation", "GERT Dispute", "GERT Iterations",
    "GERT Verification", "Warnings and Limitations", "Metadata",
]

TEAL = "0F766E"
INK = "1F2937"
WHITE = "FFFFFF"
LIGHT = "F8FAFC"
LINE = "CBD5E1"
GREEN = "15803D"
GREEN_LIGHT = "DCFCE7"
AMBER = "B45309"
AMBER_LIGHT = "FEF3C7"

HEADER_FILL = PatternFill("solid", fgColor=TEAL)
TITLE_FILL = PatternFill("solid", fgColor=INK)
LIGHT_FILL = PatternFill("solid", fgColor=LIGHT)
PASS_FILL = PatternFill("solid", fgColor=GREEN_LIGHT)
WARN_FILL = PatternFill("solid", fgColor=AMBER_LIGHT)
THIN_BOTTOM = Border(bottom=Side(style="thin", color=LINE))


def configure_sheet(ws, widths: list[float] | None = None, freeze: str | None = "A2") -> None:
    ws.sheet_view.showGridLines = False
    if freeze:
        ws.freeze_panes = freeze
    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width


def styled_row(ws, values, *, kind: str = "data"):
    cells = []
    for value in values:
        cell = WriteOnlyCell(ws, value=value)
        cell.alignment = Alignment(vertical="center", wrap_text=kind in {"header", "title", "warning"})
        if kind == "header":
            cell.fill = HEADER_FILL
            cell.font = Font(color=WHITE, bold=True)
            cell.border = THIN_BOTTOM
        elif kind == "title":
            cell.fill = TITLE_FILL
            cell.font = Font(color=WHITE, bold=True, size=16)
        elif kind == "pass":
            cell.fill = PASS_FILL
            cell.font = Font(color=GREEN, bold=True)
        elif kind == "warning":
            cell.fill = WARN_FILL
            cell.font = Font(color=AMBER, bold=True)
        cells.append(cell)
    ws.append(cells)


def append_table(ws, headers, rows) -> None:
    styled_row(ws, headers, kind="header")
    for row in rows:
        ws.append(row)


def add_bar_chart(source_ws, target_ws, title, category_col, data_cols, min_row, max_row, anchor, percent=False):
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.height = 8.2
    chart.width = 14.5
    for data_col in data_cols:
        chart.add_data(Reference(source_ws, min_col=data_col, min_row=min_row, max_row=max_row), titles_from_data=True)
    chart.set_categories(Reference(source_ws, min_col=category_col, min_row=min_row + 1, max_row=max_row))
    chart.legend = None if len(data_cols) == 1 else chart.legend
    chart.y_axis.numFmt = "0.0%" if percent else "0"
    target_ws.add_chart(chart, anchor)
    return chart


def add_line_chart(source_ws, target_ws, title, category_col, data_cols, min_row, max_row, anchor, percent=False):
    chart = LineChart()
    chart.style = 13
    chart.title = title
    chart.height = 8.2
    chart.width = 14.5
    for data_col in data_cols:
        chart.add_data(Reference(source_ws, min_col=data_col, min_row=min_row, max_row=max_row), titles_from_data=True)
    chart.set_categories(Reference(source_ws, min_col=category_col, min_row=min_row + 1, max_row=max_row))
    chart.legend = None if len(data_cols) == 1 else chart.legend
    chart.y_axis.numFmt = "0.0%" if percent else "0.00"
    target_ws.add_chart(chart, anchor)
    return chart


def audit_rows():
    for r in DATA["records"]:
        yield [
            r["word_row_number"], r["source_node_name"], r["node_id"], r["input_logic"], r["output_logic"],
            r["arc_expression"], r["arc_tag"], r["probability"], r["loop_cap"], r["o"], r["ml"], r["p"],
            r["probability_group_total"], r["missing_fields"], r["pert_use"], r["gert_use"],
            r["validation_status"], r["validation_message"],
        ]


AUDIT_HEADERS = [
    "Word table row", "Source node name", "Node ID", "Input logic", "Output logic", "Arc expression",
    "Arc Tag", "Probability", "Loop Cap", "O", "ML", "P", "Probability-group total", "Missing fields",
    "PERT use", "GERT use", "Validation status", "Validation message",
]

GERT_HEADERS = [
    "Iteration ID", "Node sequence", "Arc sequence", "Raw traversal durations", "Rounded traversal durations",
    "Total duration", "Transition count", "Loop counts", "Cap events", "Original probabilities",
    "Effective probabilities", "Renormalisation events", "Dispute visited", "Final terminal", "Final outcome",
    "Reconciliation status",
]


def gert_row(r):
    return [
        r["iteration_id"], r["node_sequence"], r["arc_sequence"], r["raw_traversal_durations"],
        r["rounded_traversal_durations"], r["total_duration"], r["transition_count"], r["loop_counts"],
        r["cap_event_count"], r["original_probabilities"], r["effective_probabilities"],
        r["renormalisation_event_count"], r["dispute_visited"], r["final_terminal"], r["final_outcome"],
        r["reconciliation_status"],
    ]


def build_main_workbook() -> None:
    wb = Workbook(write_only=True)

    ws = wb.create_sheet("Executive Summary")
    configure_sheet(ws, [32, 72, 3, 55], freeze=None)
    styled_row(ws, [DATA["title"], "", "", ""], kind="title")
    styled_row(ws, [DATA["scientific_statement"], "", "", ""], kind="warning")
    s = DATA["pert"]["summary"]
    g = DATA["gert"]["summary"]
    append_table(ws, ["Metric", "Value", "", "Critical GERT warning"], [
        ["Scientific status", "Preliminary experimental results; not final thesis findings", "", DATA["gert"]["warnings"][0]],
        ["Input file", DATA["source"]["file_name"], "", DATA["gert"]["warnings"][1]],
        ["Input SHA-256", DATA["source"]["sha256"], "", ""],
        ["PERT iterations", s["n"], "", ""],
        ["PERT simulated mean (days)", s["mean"], "", ""],
        ["PERT analytical rounded mean (days)", s["analytical_rounded_mean"], "", ""],
        ["PERT verification", s["verification_status"], "", ""],
        ["PERT reproducibility", s["reproducibility_status"], "", ""],
        ["PERT convergence", s["convergence_status"], "", ""],
        ["GERT status", DATA["gert"]["status"], "", ""],
        ["GERT successful / terminated", f"{g['successful_count']} / {g['terminated_count']}", "", ""],
        ["GERT exact verification", g["verification_status"], "", ""],
        ["Tests passed / failed", f"{DATA['test_summary']['passed']} / {DATA['test_summary']['failed']}", "", ""],
    ])

    ws = wb.create_sheet("Run Information")
    configure_sheet(ws, [34, 105])
    append_table(ws, ["Field", "Value"], [
        ["Title", DATA["title"]], ["Root seed", DATA["configuration"]["seed"]],
        ["Iterations", DATA["configuration"]["iterations"]], ["Time unit", DATA["configuration"]["time_unit"]],
        ["Beta-PERT lambda", DATA["configuration"]["beta_pert_lambda"]],
        ["Python version", DATA["configuration"]["python_version"]], ["NumPy version", DATA["configuration"]["numpy_version"]],
        ["SciPy version", DATA["configuration"]["scipy_version"]], ["Input file", DATA["source"]["path"]],
        ["Input SHA-256", DATA["source"]["sha256"]], ["Input table row count", DATA["source"]["table_row_count"]],
        ["Physical Word table columns", DATA["network"]["physical_column_count"]],
        ["Run started (UTC)", DATA["pert"]["started"]], ["Run completed (UTC)", DATA["pert"]["completed"]],
        *[[stream["name"], f"entropy={stream['entropy']}; spawn_key={'.'.join(map(str, stream['spawn_key']))}; pool_size={stream['pool_size']}"] for stream in DATA["pert"]["stream_metadata"]],
    ])

    ws = wb.create_sheet("Input Audit")
    configure_sheet(ws, [14, 40, 12, 14, 14, 18, 14, 13, 11, 9, 9, 9, 20, 28, 28, 20, 18, 54])
    append_table(ws, AUDIT_HEADERS, audit_rows())

    ws = wb.create_sheet("Word Parsed Records")
    configure_sheet(ws, [11, 12, 42, 18, 18, 12, 12, 14, 14, 12, 10, 12, 12])
    append_table(ws, ["Word row", "Node ID", "Node Name", "Node Type Input", "Node Type Output", "From", "To", "Arc Tag", "Probability", "Loop Cap", "O", "ML", "P"], ([i + 1, *row] for i, row in enumerate(DATA["raw_word_rows"])))

    ws = wb.create_sheet("Validation")
    configure_sheet(ws, [62, 18])
    styled_row(ws, ["Test", "Status"], kind="header")
    for test in DATA["tests"]:
        styled_row(ws, [test["test"], test["status"]], kind="pass")

    ws = wb.create_sheet("PERT Input")
    configure_sheet(ws, [12, 38, 10, 10, 10, 15, 15, 22, 26])
    append_table(ws, ["Sequence", "Activity", "O", "ML", "P", "Alpha", "Beta", "Pre-ceiling mean", "Analytical rounded mean"], ([r["sequence"], r["activity"], r["o"], r["ml"], r["p"], r["alpha"], r["beta"], r["pre_ceiling_mean"], r["analytical_rounded_mean"]] for r in DATA["pert"]["analytical_rows"]))

    ws = wb.create_sheet("PERT Summary")
    configure_sheet(ws, [38, 28, 24])
    append_table(ws, ["Statistic", "Value", "Unit / status"], [
        ["N", s["n"], "iterations"], ["Mean", s["mean"], "days"], ["Ceiling of Mean", s["ceiling_of_mean"], "days"],
        ["Variance", s["variance"], "days^2"], ["Standard deviation", s["standard_deviation"], "days"],
        ["Coefficient of variation", s["coefficient_of_variation"], "ratio"], ["Minimum", s["minimum"], "days"],
        ["Maximum", s["maximum"], "days"], ["P25", s["p25"], "days"], ["P50", s["p50"], "days"],
        ["P75", s["p75"], "days"], ["IQR", s["iqr"], "days"], ["P80", s["p80"], "days"],
        ["P90", s["p90"], "days"], ["P95", s["p95"], "days"], ["P99", s["p99"], "days"],
        ["Monte Carlo standard error", s["mcse"], "days"], ["Analytical rounded mean", s["analytical_rounded_mean"], "days"],
        ["Simulated - analytical", s["difference"], "days"], ["Monte Carlo verification", s["verification_status"], ""],
        ["Convergence", s["convergence_status"], ""], ["Reproducibility", s["reproducibility_status"], ""],
    ])

    ws = wb.create_sheet("PERT Analytical Check")
    configure_sheet(ws, [12, 38, 15, 15, 22, 26, 20, 18])
    analytical_rows = [[r["sequence"], r["activity"], r["alpha"], r["beta"], r["pre_ceiling_mean"], r["analytical_rounded_mean"], r["simulated_mean"], r["simulated_mean"] - r["analytical_rounded_mean"]] for r in DATA["pert"]["analytical_rows"]]
    analytical_rows.append(["TOTAL", "Six-activity route", "", "", DATA["pert_gates"]["analytical_pre_ceiling_mean"], s["analytical_rounded_mean"], s["mean"], s["difference"]])
    append_table(ws, ["Sequence", "Activity", "Alpha", "Beta", "Pre-ceiling mean", "Analytical rounded mean", "Simulated mean", "Difference"], analytical_rows)

    ws = wb.create_sheet("PERT P1-P99")
    configure_sheet(ws, [18, 32])
    append_table(ws, ["Percentile", "Nearest-rank duration (days)"], ([p, DATA["pert"]["percentiles"][str(p)]] for p in range(1, 100)))

    ws = wb.create_sheet("PERT Iterations")
    configure_sheet(ws, [18, 26])
    append_table(ws, ["Iteration ID", "Total duration (days)"], ([i + 1, value] for i, value in enumerate(DATA["pert"]["totals"])))

    ws = wb.create_sheet("PERT Activity Durations")
    configure_sheet(ws, [18, 20, 38, 24, 28])
    styled_row(ws, ["Iteration ID", "Activity sequence", "Activity", "Raw duration (days)", "Rounded duration (days)"], kind="header")
    for i in range(DATA["configuration"]["iterations"]):
        for a, activity in enumerate(DATA["activities"]):
            ws.append([i + 1, a + 1, activity["activity"], DATA["pert"]["raw"][i][a], DATA["pert"]["rounded"][i][a]])

    ws = wb.create_sheet("PERT Convergence")
    configure_sheet(ws, [16, 18, 18, 12, 12, 12, 12, 12, 26, 16])
    append_table(ws, ["Sample size", "Running mean", "Running SD", "P50", "P80", "P90", "P95", "P99", "Maximum relative change", "Stability"], ([r["sample_size"], r["running_mean"], r["running_sd"], r["p50"], r["p80"], r["p90"], r["p95"], r["p99"], r["max_relative_change"], "N/A" if r["stability_pass"] is None else "PASS" if r["stability_pass"] else "FAIL"] for r in DATA["pert"]["convergence_rows"]))

    ws = wb.create_sheet("PERT Histogram Data")
    configure_sheet(ws, [18, 16, 22])
    append_table(ws, ["Duration (days)", "Frequency", "Relative frequency"], ([r["duration"], r["frequency"], r["relative_frequency"]] for r in DATA["pert"]["histogram"]))

    ws = wb.create_sheet("PERT ECDF Data")
    configure_sheet(ws, [18, 26])
    append_table(ws, ["Duration (days)", "Cumulative probability"], ([r["duration"], r["cumulative_probability"]] for r in DATA["pert"]["ecdf"]))

    chart_data = wb.create_sheet("PERT Chart Data")
    configure_sheet(chart_data, [14, 14, 20, 3, 14, 24, 3, 14, 14, 3, 16, 18, 18, 12, 12, 12, 12, 12, 3, 36, 24])
    headers = ["Duration", "Frequency", "Relative frequency", "", "Duration", "Cumulative probability", "", "Percentile", "Duration", "", "Sample size", "Running mean", "Running SD", "P50", "P80", "P90", "P95", "P99", "", "Activity", "Mean rounded duration"]
    styled_row(chart_data, headers, kind="header")
    max_rows = max(len(DATA["pert"]["histogram"]), len(DATA["pert"]["ecdf"]), 99, len(DATA["pert"]["convergence_rows"]), 6)
    for i in range(max_rows):
        row = [None] * 21
        if i < len(DATA["pert"]["histogram"]):
            h = DATA["pert"]["histogram"][i]
            row[0:3] = [h["duration"], h["frequency"], h["relative_frequency"]]
        if i < len(DATA["pert"]["ecdf"]):
            e = DATA["pert"]["ecdf"][i]
            row[4:6] = [e["duration"], e["cumulative_probability"]]
        if i < 99:
            row[7:9] = [i + 1, DATA["pert"]["percentiles"][str(i + 1)]]
        if i < len(DATA["pert"]["convergence_rows"]):
            c = DATA["pert"]["convergence_rows"][i]
            row[10:18] = [c["sample_size"], c["running_mean"], c["running_sd"], c["p50"], c["p80"], c["p90"], c["p95"], c["p99"]]
        if i < 6:
            row[19:21] = [DATA["activities"][i]["activity"], DATA["pert"]["analytical_rows"][i]["simulated_mean"]]
        chart_data.append(row)

    charts_ws = wb.create_sheet("PERT Charts")
    configure_sheet(charts_ws, [12] * 17, freeze=None)
    styled_row(charts_ws, ["PERT Monte Carlo Charts"], kind="title")
    add_bar_chart(chart_data, charts_ws, "PERT total duration histogram", 1, [2], 1, len(DATA["pert"]["histogram"]) + 1, "A3")
    add_line_chart(chart_data, charts_ws, "PERT empirical CDF", 5, [6], 1, len(DATA["pert"]["ecdf"]) + 1, "J3", percent=True)
    add_line_chart(chart_data, charts_ws, "PERT P1-P99 curve", 8, [9], 1, 100, "A20")
    add_line_chart(chart_data, charts_ws, "PERT mean convergence", 11, [12], 1, len(DATA["pert"]["convergence_rows"]) + 1, "J20")
    add_line_chart(chart_data, charts_ws, "PERT percentile convergence", 11, [14, 15, 16, 17, 18], 1, len(DATA["pert"]["convergence_rows"]) + 1, "A37")
    add_bar_chart(chart_data, charts_ws, "Mean activity duration", 20, [21], 1, 7, "J37")

    ws = wb.create_sheet("GERT Status")
    configure_sheet(ws, [30, 80, 3, 90], freeze=None)
    styled_row(ws, ["GERT Simulation Status", "", "", ""], kind="title")
    append_table(ws, ["Field", "Value", "", "Warning"], [
        ["Status", DATA["gert"]["status"], "", DATA["gert"]["warnings"][0]],
        ["Structural gate", DATA["gert"]["structural_gate"], "", DATA["gert"]["warnings"][1]],
        ["Iterations", g["valid_runs"], "", ""], ["Exact verification", g["verification_status"], "", ""],
        ["Reproducibility", g["reproducibility_status"], "", ""], ["Reason when not run", DATA["gert"]["reason"], "", ""],
    ])

    ws = wb.create_sheet("GERT Input Audit")
    configure_sheet(ws, [12, 15, 40, 14, 12, 10, 10, 10, 12, 18, 14, 12, 12, 16, 14, 18, 24, 16])
    append_table(ws, ["Record ID", "Word row number", "Source Node Name", "Probability", "Loop Cap", "O", "ML", "P", "Node ID", "Arc expression", "Arc Tag", "From Node", "To Node", "Is Return Arc", "Loop ID", "Terminal Outcome", "Missing Fields", "Status"], ([r["record_id"], r["word_row_number"], r["source_node_name"], r["probability"], r["loop_cap"], r["o"], r["ml"], r["p"], r["node_id"], r["arc_expression"], r["arc_tag"], r["from_node"], r["to_node"], r["is_return_arc"], r["loop_id"], r["terminal_outcome"], r["missing_fields"], "SIMULATED"] for r in DATA["records"]))

    outcome_ws = wb.create_sheet("GERT Outcome Summary")
    configure_sheet(outcome_ws, [34, 24, 3, 18, 18, 3, 20, 24, 3, 20, 18, 3, 20, 18, 3] + [12] * 16)
    outcome_metrics = [
        ["Valid runs", g["valid_runs"]], ["Invalid runs", g["invalid_runs"]], ["Successful count", g["successful_count"]],
        ["Terminated count", g["terminated_count"]], ["Successful probability", g["successful_probability"]],
        ["Terminated probability", g["terminated_probability"]], ["Overall mean duration", g["overall_duration"]["mean"]],
        ["Overall SD", g["overall_duration"]["standard_deviation"]], ["Overall P50", g["overall_duration"]["p50"]],
        ["Overall P95", g["overall_duration"]["p95"]], ["Dispute visited", g["dispute_visited_count"]],
        ["Cap events", g["cap_event_count"]], ["Renormalisation events", g["renormalisation_event_count"]],
        ["Verification", g["verification_status"]],
    ]
    term_ecdf = DATA["gert"]["outcome_ecdf"]["Terminated"]
    success_hist = DATA["gert"]["histograms"]["Successful"]
    term_hist = DATA["gert"]["histograms"]["Terminated"]
    styled_row(outcome_ws, ["Metric", "Value", "", "Outcome", "Probability", "", "Terminated duration", "Cumulative probability", "", "Successful duration", "Frequency", "", "Terminated duration", "Frequency"], kind="header")
    count = max(len(outcome_metrics), 2, len(term_ecdf), len(success_hist), len(term_hist))
    for i in range(count):
        row = [None] * 14
        if i < len(outcome_metrics): row[0:2] = outcome_metrics[i]
        if i < 2: row[3:5] = [["Successful", g["successful_probability"]], ["Terminated", g["terminated_probability"]]][i]
        if i < len(term_ecdf): row[6:8] = [term_ecdf[i]["duration"], term_ecdf[i]["cumulative_probability"]]
        if i < len(success_hist): row[9:11] = [success_hist[i]["duration"], success_hist[i]["frequency"]]
        if i < len(term_hist): row[12:14] = [term_hist[i]["duration"], term_hist[i]["frequency"]]
        outcome_ws.append(row)
    add_bar_chart(outcome_ws, outcome_ws, "GERT outcome probabilities", 4, [5], 1, 3, "P2", percent=True)
    add_bar_chart(outcome_ws, outcome_ws, "Successful duration histogram (no observations)", 10, [11], 1, len(success_hist) + 1, "X2")
    add_bar_chart(outcome_ws, outcome_ws, "Terminated duration histogram", 13, [14], 1, len(term_hist) + 1, "P19")
    add_line_chart(outcome_ws, outcome_ws, "GERT outcome ECDF", 7, [8], 1, len(term_ecdf) + 1, "X19", percent=True)

    ws = wb.create_sheet("GERT Successful Results")
    configure_sheet(ws, [16, 56, 40, 42, 42, 18, 18, 48, 14, 38, 38, 22, 18, 16, 18, 20])
    styled_row(ws, GERT_HEADERS, kind="header")
    if DATA["gert"]["successful_results"]:
        for r in DATA["gert"]["successful_results"]: ws.append(gert_row(r))
    else:
        ws.append(["No successful iterations observed"] + [None] * 15)

    ws = wb.create_sheet("GERT Terminated Results")
    configure_sheet(ws, [16, 56, 40, 42, 42, 18, 18, 48, 14, 38, 38, 22, 18, 16, 18, 20])
    append_table(ws, GERT_HEADERS, (gert_row(r) for r in DATA["gert"]["terminated_results"]))

    paths_ws = wb.create_sheet("GERT Paths")
    configure_sheet(paths_ws, [10, 72, 56, 18, 16, 16, 3, 14, 16] + [12] * 10)
    styled_row(paths_ws, ["Rank", "Node sequence", "Arc sequence", "Outcome", "Frequency", "Probability", "", "Path rank", "Frequency"], kind="header")
    top_paths = DATA["gert"]["paths"][:10]
    for i, r in enumerate(DATA["gert"]["paths"]):
        row = [r["rank"], r["node_sequence"], r["arc_sequence"], r["outcome"], r["frequency"], r["probability"], None, None, None]
        if i < len(top_paths): row[7:9] = [f"Path {top_paths[i]['rank']}", top_paths[i]["frequency"]]
        paths_ws.append(row)
    add_bar_chart(paths_ws, paths_ws, "Top realised GERT paths", 8, [9], 1, len(top_paths) + 1, "K2")

    ws = wb.create_sheet("GERT Nodes")
    configure_sheet(ws, [16, 18, 28])
    append_table(ws, ["Node ID", "Visit count", "Mean visits per iteration"], ([r["node_id"], r["visit_count"], r["mean_visits_per_iteration"]] for r in DATA["gert"]["nodes"]))

    ws = wb.create_sheet("GERT Arcs")
    configure_sheet(ws, [14, 14, 14, 22, 18, 30])
    append_table(ws, ["Arc Tag", "From Node", "To Node", "Original probability", "Traversal count", "Mean traversals per iteration"], ([r["arc_tag"], r["from_node"], r["to_node"], r["original_probability"], r["traversal_count"], r["mean_traversals_per_iteration"]] for r in DATA["gert"]["arcs"]))

    loops_ws = wb.create_sheet("GERT Loops")
    configure_sheet(loops_ws, [14, 14, 14, 14, 12, 18, 22, 18, 20, 3, 14, 18, 3, 18, 16] + [12] * 10)
    styled_row(loops_ws, ["Loop ID", "Arc Tag", "From Node", "To Node", "Loop Cap", "Activation count", "Activation probability", "Total traversals", "Maximum observed", "", "Loop", "Activation count", "", "Total repetitions", "Frequency"], kind="header")
    dist = DATA["gert"]["loop_repetition_distribution"]
    count = max(len(DATA["gert"]["loops"]), len(dist))
    for i in range(count):
        row = [None] * 15
        if i < len(DATA["gert"]["loops"]):
            r = DATA["gert"]["loops"][i]
            row[0:9] = [r["loop_id"], r["arc_tag"], r["from_node"], r["to_node"], r["loop_cap"], r["activation_count"], r["activation_probability"], r["total_traversals"], r["maximum_observed"]]
            row[10:12] = [r["arc_tag"], r["activation_count"]]
        if i < len(dist): row[13:15] = [dist[i]["repetitions"], dist[i]["frequency"]]
        loops_ws.append(row)
    add_bar_chart(loops_ws, loops_ws, "GERT loop activation", 11, [12], 1, len(DATA["gert"]["loops"]) + 1, "Q2")
    add_bar_chart(loops_ws, loops_ws, "GERT loop repetition distribution", 14, [15], 1, len(dist) + 1, "Q20")

    ws = wb.create_sheet("GERT Cap Events")
    configure_sheet(ws, [24, 18, 18, 12, 24])
    rows = ([r["iteration_id"], r["loop_id"], r["arc_tag"], r["cap"], r["count_when_reached"]] for r in DATA["gert"]["cap_events"])
    append_table(ws, ["Iteration ID", "Loop ID", "Arc Tag", "Cap", "Count when reached"], rows)
    if not DATA["gert"]["cap_events"]: ws.append(["No cap events observed"])

    ws = wb.create_sheet("GERT Renormalisation")
    configure_sheet(ws, [30, 16, 32, 40, 42, 42])
    append_table(ws, ["Iteration ID", "Node", "Removed Arc Tags", "Eligible Arc Tags", "Original probabilities", "Effective probabilities"], ([r["iteration_id"], r["node"], r["removed_arc_tags"], r["eligible_arc_tags"], r["original_probabilities"], r["effective_probabilities"]] for r in DATA["gert"]["renormalisation_events"]))
    if not DATA["gert"]["renormalisation_events"]: ws.append(["No renormalisation events observed"])

    ws = wb.create_sheet("GERT Dispute")
    configure_sheet(ws, [28, 30])
    append_table(ws, ["Metric", "Value"], [["Dispute node ID", "SD"], ["Dispute visited count", g["dispute_visited_count"]], ["Dispute terminal count", 0], ["Classification", "Transient node"]])

    ws = wb.create_sheet("GERT Iterations")
    configure_sheet(ws, [16, 56, 40, 42, 42, 18, 18, 48, 14, 38, 38, 22, 18, 16, 18, 20])
    append_table(ws, GERT_HEADERS, (gert_row(r) for r in DATA["gert"]["iterations"]))

    ws = wb.create_sheet("GERT Verification")
    configure_sheet(ws, [46, 24, 42])
    exact = DATA["gert"]["exact"]
    rows = [[r["check"], r["status"], "Monte Carlo / run-level check"] for r in DATA["gert"]["verification_checks"]]
    rows.extend([
        ["Reachable transient states", exact["reachable_transient_states"], "Independent finite-state verifier"],
        ["Exact successful probability", exact["successful_probability"], "Independent finite-state verifier"],
        ["Exact terminated probability", exact["terminated_probability"], "Independent finite-state verifier"],
        ["Exact probability sum", exact["probability_sum"], "Independent finite-state verifier"],
        ["Exact expected rounded duration", exact["expected_rounded_duration"], "Independent finite-state verifier"],
        ["Overall verification", g["verification_status"], "Required acceptance gate"],
    ])
    append_table(ws, ["Check", "Result", "Method"], rows)

    ws = wb.create_sheet("Warnings and Limitations")
    configure_sheet(ws, [30, 120])
    append_table(ws, ["Category", "Statement"], [
        ["Scientific status", "These are PRELIMINARY EXPERIMENTAL SIMULATION RESULTS and do not constitute final thesis findings."],
        ["Input authority", "All numerical and routing values were taken from the most recently attached Word file."],
        ["Word table structure", "The latest file contains 12 physical columns: From and To are separate fields within the outgoing-arc definition."],
        ["GERT routing warning", DATA["gert"]["warnings"][0]], ["GERT reachability warning", DATA["gert"]["warnings"][1]],
        ["No silent inference", "The Arc Tag e12 was not used to override the explicit To Node value S1."],
        ["Approval limitation", "Preliminary results do not constitute final thesis findings until the network structure and expert inputs are formally approved."],
    ])

    ws = wb.create_sheet("Metadata")
    configure_sheet(ws, [32, 120])
    append_table(ws, ["Field", "Value"], [
        ["Title", DATA["title"]], ["Input file", DATA["source"]["path"]], ["Input SHA-256", DATA["source"]["sha256"]],
        ["Parsed nodes", "; ".join(DATA["nodes_seen"])], ["Probabilistic records", len(DATA["records"])],
        ["Loop Cap = 2 records", sum(r["loop_cap"] == 2 for r in DATA["records"])], ["Root seed", DATA["configuration"]["seed"]],
        ["PERT stream", json.dumps(DATA["pert"]["stream_metadata"][0])], ["GERT routing stream", json.dumps(DATA["pert"]["stream_metadata"][1])],
        ["GERT duration stream", json.dumps(DATA["pert"]["stream_metadata"][2])], ["Python version", DATA["configuration"]["python_version"]],
        ["NumPy version", DATA["configuration"]["numpy_version"]], ["SciPy version", DATA["configuration"]["scipy_version"]],
        ["Main workbook", str(MAIN_PATH)], ["Input audit workbook", str(AUDIT_PATH)], ["Report", str(REPORT_PATH)],
    ])

    if wb.sheetnames != REQUIRED_SHEETS:
        raise AssertionError("Workbook sheet order does not match the required 32-sheet list.")
    wb.save(MAIN_PATH)


def build_audit_workbook() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Input Audit"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.append(AUDIT_HEADERS)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color=WHITE, bold=True)
    for row in audit_rows():
        ws.append(row)
    for index, width in enumerate([14, 40, 12, 14, 14, 18, 14, 13, 11, 9, 9, 9, 20, 28, 28, 20, 18, 54], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    validation = wb.create_sheet("Validation")
    validation.append(["Check", "Result"])
    for cell in validation[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color=WHITE, bold=True)
    for row in [
        ["Parsed node groups", len(DATA["nodes_seen"])], ["Probabilistic records", len(DATA["records"])],
        ["Probability groups", 7], ["Probability group sums", "PASS"], ["Loop Cap = 2 records", 10],
        ["GERT structural gate", DATA["gert"]["structural_gate"]], ["GERT status", DATA["gert"]["status"]],
    ]:
        validation.append(row)
    metadata = wb.create_sheet("Metadata")
    metadata.append(["Field", "Value"])
    for cell in metadata[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color=WHITE, bold=True)
    metadata.append(["Input file", DATA["source"]["path"]])
    metadata.append(["SHA-256", DATA["source"]["sha256"]])
    metadata.append(["Generated title", DATA["title"]])
    wb.save(AUDIT_PATH)


def write_report() -> None:
    s = DATA["pert"]["summary"]
    g = DATA["gert"]["summary"]
    report = f"""# {DATA['title']}

{DATA['scientific_statement']}

## A. Input-file audit

- File name: {DATA['source']['file_name']}
- SHA-256: {DATA['source']['sha256']}
- First-table rows: {DATA['source']['table_row_count']}
- Physical columns: {DATA['network']['physical_column_count']}
- Parsed nodes: {'; '.join(DATA['nodes_seen'])}
- Parsed probabilistic records: {len(DATA['records'])}
- Probability-group validation: PASS; all seven groups sum to 1.00 within 1e-12
- Loop Cap = 2 records: {sum(r['loop_cap'] == 2 for r in DATA['records'])}
- Missing structural fields: none in the latest file

## B. PERT results

- N: {s['n']}
- Seed: {DATA['configuration']['seed']}
- Activity count: {DATA['pert_gates']['activity_count']}
- Sum O / ML / P: {DATA['pert_gates']['sum_o']} / {DATA['pert_gates']['sum_ml']} / {DATA['pert_gates']['sum_p']} days
- Analytical pre-ceiling mean: {DATA['pert_gates']['analytical_pre_ceiling_mean']:.10f} days
- Analytical rounded mean: {s['analytical_rounded_mean']:.12f} days
- Simulated mean: {s['mean']:.6f} days
- Monte Carlo standard error: {s['mcse']:.12f} days
- Simulated - analytical: {s['difference']:.12f} days
- Minimum / maximum: {s['minimum']} / {s['maximum']} days
- P50 / P80 / P90 / P95 / P99: {s['p50']} / {s['p80']} / {s['p90']} / {s['p95']} / {s['p99']} days
- Convergence: {s['convergence_status']}
- Reproducibility: {s['reproducibility_status']}
- Validation: {s['verification_status']}

## C. GERT status

- Status: {DATA['gert']['status']}
- Structural gate: {DATA['gert']['structural_gate']}
- Valid / invalid runs: {g['valid_runs']} / {g['invalid_runs']}
- Successful / terminated outcomes: {g['successful_count']} / {g['terminated_count']}
- Successful / terminated probabilities: {g['successful_probability']:.4%} / {g['terminated_probability']:.4%}
- Exact finite-state verification: {g['verification_status']}
- Reproducibility: {g['reproducibility_status']}
- Routing warning: {DATA['gert']['warnings'][0]}
- Reachability warning: {DATA['gert']['warnings'][1]}

## D. Limitations

The simulation follows the explicit Word fields exactly. In particular, the e12 record is simulated as S1 to S1; its tag is not used to infer S1 to S2. This explicit mapping makes Final Closeout unreachable from the start state and is the reason all GERT outcomes terminate.

Preliminary results do not constitute final thesis findings until the network structure and expert inputs are formally approved.
"""
    REPORT_PATH.write_text(report, encoding="utf-8")


def verify_outputs() -> dict:
    reopened = load_workbook(MAIN_PATH, read_only=True, data_only=False)
    if reopened.sheetnames != REQUIRED_SHEETS:
        raise AssertionError("Reopened workbook does not contain the required sheets in order.")
    required_nonempty = ["Input Audit", "PERT Iterations", "PERT Activity Durations", "PERT Charts", "GERT Status", "GERT Input Audit", "GERT Outcome Summary", "GERT Iterations", "GERT Verification"]
    dimensions = {name: reopened[name].calculate_dimension(force=True) for name in required_nonempty}
    reopened.close()

    with zipfile.ZipFile(MAIN_PATH) as archive:
        names = archive.namelist()
        chart_files = [name for name in names if name.startswith("xl/charts/chart") and name.endswith(".xml")]
        if len(chart_files) != 13:
            raise AssertionError(f"Expected 13 native charts, found {len(chart_files)}.")
        external_entries = [name for name in names if name.startswith("xl/externalLinks/")]
        if external_entries:
            raise AssertionError(f"External-link parts found: {external_entries}")
        relation_text = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in names
            if name.endswith(".rels")
        )
        if 'TargetMode="External"' in relation_text:
            raise AssertionError("External relationship found in workbook.")
        chart_text = "\n".join(archive.read(name).decode("utf-8", errors="ignore") for name in chart_files)
        if "[" in chart_text or "]" in chart_text:
            raise AssertionError("A chart series contains an external workbook reference.")
        if not all(re.search(r"<(?:c:)?ser>", archive.read(name).decode("utf-8", errors="ignore")) for name in chart_files):
            raise AssertionError("At least one native chart contains no series.")

    audit_reopened = load_workbook(AUDIT_PATH, read_only=True)
    if audit_reopened["Input Audit"].max_row != 30:
        raise AssertionError("Preliminary_Input_Audit.xlsx does not contain all 29 records.")
    audit_reopened.close()
    if not REPORT_PATH.exists() or REPORT_PATH.stat().st_size == 0:
        raise AssertionError("Markdown report is missing or empty.")
    return {
        "required_sheets": len(REQUIRED_SHEETS),
        "native_charts": len(chart_files),
        "external_links": 0,
        "dimensions": dimensions,
        "tests_passed": DATA["test_summary"]["passed"],
        "tests_failed": DATA["test_summary"]["failed"],
    }


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_main_workbook()
    build_audit_workbook()
    write_report()
    verification = verify_outputs()
    (WORK_DIR / "final_verification.json").write_text(json.dumps(verification, indent=2), encoding="utf-8")
    print(json.dumps({
        "main": str(MAIN_PATH), "report": str(REPORT_PATH), "audit": str(AUDIT_PATH),
        "verification": verification,
    }, indent=2))


if __name__ == "__main__":
    main()
