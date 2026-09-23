from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Dict, Iterable, List

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

from .engine import Case, asdict, run_gert, run_pert, validate_case


MASTER_SHEETS = [
    "Executive Summary", "Run Information", "Validation Results", "PERT Input Snapshot", "GERT Input Snapshot",
    "Network Validation", "PERT Summary", "PERT P1-P99", "PERT Activity Results", "PERT Convergence",
    "GERT Outcome Summary", "GERT Successful Summary", "GERT Successful P1-P99", "GERT Terminated Summary",
    "GERT Terminated P1-P99", "GERT Path Analysis", "GERT Node Analysis", "GERT Arc Analysis", "GERT Loop Summary",
    "Loop Repetition Distribution", "Loop Active-Inactive Comparison", "Loop Cap Outcomes", "Renormalisation Summary",
    "Renormalisation Events", "Dispute Analysis", "Transition Count Analysis", "GERT Convergence",
    "PERT-GERT Comparison", "Invalid Iterations", "Metadata"
]


def rows_to_sheet(ws, rows: List[Dict], empty_message: str = "No rows") -> None:
    if not rows:
        ws.append([empty_message])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(42, max(12, max(len(str(c.value or "")) for c in col) + 2))


def dict_to_sheet(ws, data: Dict) -> None:
    ws.append(["Metric", "Value"])
    for k, v in data.items():
        ws.append([k, v])


def add_chart(ws, title: str, data_col: int = 2) -> None:
    if ws.max_row < 3 or ws.max_column < data_col:
        return
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = "Value"
    chart.x_axis.title = "Index"
    data = Reference(ws, min_col=data_col, min_row=1, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    ws.add_chart(chart, "E2")


def write_workbook(path: Path, sheets: Dict[str, object]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, data in sheets.items():
        ws = wb.create_sheet(name[:31])
        if isinstance(data, list):
            rows_to_sheet(ws, data)
        elif isinstance(data, dict):
            dict_to_sheet(ws, data)
        else:
            ws.append([str(data)])
    wb.save(path)


def export_csv(path: Path, rows: List[Dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def generate_outputs(case: Case, root: Path, iterations: int) -> Dict[str, Path]:
    root.mkdir(parents=True, exist_ok=True)
    raw = root / "Raw_Data"
    charts = root / "Charts_Excel"
    chart_img = root / "Charts_PNG"
    charts.mkdir(exist_ok=True)
    chart_img.mkdir(exist_ok=True)
    errors = validate_case(case)
    if errors:
        write_workbook(root / "01_Simulation_Results_Master.xlsx", {"Validation Results": [{"Error": e} for e in errors]})
        return {"status": root / "01_Simulation_Results_Master.xlsx"}
    pert = run_pert(case, iterations)
    gert = run_gert(case, iterations)
    pert_rows = [{"Iteration": i + 1, "Duration": v} for i, v in enumerate(pert["totals"])]
    comparison = compare(pert["summary"], gert["successful_summary"])
    sheets = {name: [] for name in MASTER_SHEETS}
    sheets.update({
        "Executive Summary": [{"Metric": "Iterations", "Value": iterations}, {"Metric": "Seed", "Value": case.seed}],
        "Run Information": [{"Field": k, "Value": v} for k, v in asdict(case).items() if k not in {"pert", "nodes", "arcs"}],
        "Validation Results": [{"Status": "Passed"}],
        "PERT Input Snapshot": [asdict(a) for a in case.pert],
        "GERT Input Snapshot": [asdict(a) for a in case.arcs],
        "PERT Summary": pert["summary"],
        "PERT P1-P99": pert["percentiles"],
        "PERT Activity Results": pert["activity"],
        "PERT Convergence": pert["convergence"],
        "GERT Outcome Summary": gert["outcome_summary"],
        "GERT Successful Summary": gert["successful_summary"],
        "GERT Successful P1-P99": gert["successful_percentiles"],
        "GERT Terminated Summary": gert["terminated_summary"],
        "GERT Terminated P1-P99": gert["terminated_percentiles"],
        "GERT Path Analysis": aggregate_paths(gert["paths"]),
        "GERT Node Analysis": gert["node_analysis"],
        "GERT Arc Analysis": gert["arc_analysis"],
        "GERT Loop Summary": gert["loop_summary"],
        "Loop Cap Outcomes": gert["cap_events"],
        "Renormalisation Events": gert["renorm_events"],
        "Dispute Analysis": gert["dispute"],
        "Transition Count Analysis": gert["transition_counts"],
        "GERT Convergence": gert["convergence"],
        "PERT-GERT Comparison": comparison,
        "Invalid Iterations": gert["invalid"],
        "Metadata": [{"Key": "Runtime", "Value": "Bundled Python source package"}],
    })
    write_workbook(root / "01_Simulation_Results_Master.xlsx", sheets)
    write_workbook(root / "02_PERT_MCS_Results.xlsx", {
        "PERT Input Snapshot": [asdict(a) for a in case.pert], "PERT Summary": pert["summary"], "P1-P99": pert["percentiles"],
        "Activity Results": pert["activity"], "Raw Iteration Totals": pert_rows, "Raw Activity Durations": pert["raw_activity"],
        "Convergence": pert["convergence"],
    })
    write_workbook(root / "03_GERT_MCS_Results.xlsx", {
        "GERT Input Snapshot": [asdict(a) for a in case.arcs], "Outcome Summary": gert["outcome_summary"],
        "Successful Summary": gert["successful_summary"], "Successful P1-P99": gert["successful_percentiles"],
        "Terminated Summary": gert["terminated_summary"], "Terminated P1-P99": gert["terminated_percentiles"],
        "Paths": aggregate_paths(gert["paths"]), "Nodes": gert["node_analysis"], "Arcs": gert["arc_analysis"],
        "Loops": gert["loop_summary"], "Loop Caps": gert["cap_events"], "Renormalisation": gert["renorm_events"],
        "Dispute": gert["dispute"], "Transition Counts": gert["transition_counts"], "Invalid": gert["invalid"],
        "Convergence": gert["convergence"],
    })
    write_workbook(root / "04_PERT_GERT_Comparison.xlsx", {"Comparison": comparison, "PERT P1-P99": pert["percentiles"], "GERT Successful P1-P99": gert["successful_percentiles"]})
    write_workbook(root / "05_Convergence_Analysis.xlsx", {"PERT": pert["convergence"], "GERT": gert["convergence"]})
    write_workbook(root / "06_Loops_Caps_Dispute_Analysis.xlsx", {"Loops": gert["loop_summary"], "Caps": gert["cap_events"], "Renormalisation": gert["renorm_events"], "Dispute": gert["dispute"]})
    write_workbook(root / "07_Raw_Simulation_Data.xlsx", {"Index": [
        {"File": "Raw_Data/PERT_Iterations.csv", "Description": "PERT iteration totals"},
        {"File": "Raw_Data/GERT_Iterations.csv", "Description": "GERT iteration outcomes"},
    ]})
    export_csv(raw / "PERT_Iterations.csv", pert_rows)
    export_csv(raw / "PERT_Activity_Durations.csv", pert["raw_activity"])
    export_csv(raw / "GERT_Iterations.csv", gert["iterations"])
    export_csv(raw / "GERT_Paths.csv", gert["paths"])
    export_csv(raw / "GERT_Nodes.csv", gert["node_analysis"])
    export_csv(raw / "GERT_Arcs.csv", gert["arc_analysis"])
    export_csv(raw / "GERT_Loops.csv", gert["loop_summary"])
    export_csv(raw / "GERT_Loop_Cap_Events.csv", gert["cap_events"])
    export_csv(raw / "GERT_Renormalisation_Events.csv", gert["renorm_events"])
    export_csv(raw / "GERT_Dispute_Results.csv", gert["dispute"])
    export_csv(raw / "GERT_Invalid_Iterations.csv", gert["invalid"])
    export_csv(raw / "Convergence_Checkpoints.csv", pert["convergence"] + gert["convergence"])
    make_chart_workbooks(charts, pert, gert, comparison, case, iterations)
    return {"status": root / "01_Simulation_Results_Master.xlsx"}


def compare(pert: Dict, gert: Dict) -> List[Dict]:
    rows = []
    for k in ["Mean", "Standard deviation", "Variance", "Coefficient of variation", "Minimum", "Maximum", "IQR", "Median"]:
        p = pert.get(k, 0)
        g = gert.get(k, 0)
        rows.append({"Metric": k, "PERT": p, "GERT Successful": g, "Absolute Difference": g - p, "Relative Difference": (g - p) / p if p else 0})
    return rows


def aggregate_paths(rows: List[Dict]) -> List[Dict]:
    counts: Dict[str, Dict] = {}
    for r in rows:
        key = (r["Path"], r["Outcome"])
        rec = counts.setdefault(key, {"Path": r["Path"], "Outcome": r["Outcome"], "Count": 0, "Total Duration": 0})
        rec["Count"] += 1
        rec["Total Duration"] += r["Duration"]
    out = []
    total = len(rows) or 1
    for rec in counts.values():
        rec["Probability"] = rec["Count"] / total
        rec["Mean Duration"] = rec["Total Duration"] / rec["Count"]
        out.append(rec)
    return sorted(out, key=lambda x: x["Count"], reverse=True)


def make_chart_workbooks(folder: Path, pert: Dict, gert: Dict, comparison: List[Dict], case: Case, iterations: int) -> None:
    chart_defs = [
        ("Chart_01_PERT_Histogram.xlsx", [{"Duration": v} for v in pert["totals"][:10000]]),
        ("Chart_02_PERT_ECDF.xlsx", pert["percentiles"]),
        ("Chart_03_PERT_P1_P99.xlsx", pert["percentiles"]),
        ("Chart_04_PERT_Mean_Convergence.xlsx", pert["convergence"]),
        ("Chart_09_GERT_Final_Outcome_Probabilities.xlsx", gert["outcome_summary"]),
        ("Chart_10_GERT_Successful_Duration_Histogram.xlsx", [{"Duration": v} for v in gert["success"][:10000]]),
        ("Chart_11_GERT_Terminated_Duration_Histogram.xlsx", [{"Duration": v} for v in gert["terminated"][:10000]]),
        ("Chart_19_Node_Visit_Probability.xlsx", gert["node_analysis"]),
        ("Chart_21_Arc_Activation_Probability.xlsx", gert["arc_analysis"]),
        ("Chart_27_Loop_Activation_Probability.xlsx", gert["loop_summary"]),
        ("Chart_49_PERT_vs_GERT_Successful.xlsx", comparison),
    ]
    for name, rows in chart_defs:
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        rows_to_sheet(ws, rows)
        cws = wb.create_sheet("Chart")
        cws.append(["Simulation", case.simulation_name])
        cws.append(["Iterations", iterations])
        cws.append(["Seed", case.seed])
        wb.save(folder / name)

