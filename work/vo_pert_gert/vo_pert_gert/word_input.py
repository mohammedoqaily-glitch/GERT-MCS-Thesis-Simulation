from __future__ import annotations

import json
import re
import hashlib
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import List, Optional

from docx import Document
from openpyxl import Workbook


EXPECTED_COLUMNS = [
    "Node ID",
    "Node Name",
    "Node Type - Input",
    "Node Type - Output",
    "Arc",
    "Tag",
    "Probability",
    "Loop Cap",
    "O",
    "ML",
    "P",
]


@dataclass
class AuditRow:
    word_table_row: int
    node_id: str
    node_name: str
    input_logic: str
    output_logic: str
    arc_expression: str
    arc_tag: str
    from_node: str
    to_node: str
    original_probability: str
    loop_cap: str
    o: str
    ml: str
    p: str
    loop_return_classification: str
    pert_route_classification: str
    validation_status: str
    source_cell_reference: str


def clean(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("\n", " | ")).strip()


def parse_arc(expr: str) -> tuple[str, str]:
    if "→" in expr:
        left, right = expr.split("→", 1)
    elif "->" in expr:
        left, right = expr.split("->", 1)
    else:
        return "", ""
    return clean(left), clean(right)


def parse_word_panel(path: Path) -> tuple[List[AuditRow], list[str]]:
    doc = Document(str(path))
    if not doc.tables:
        return [], ["No Word tables found."]
    table = doc.tables[0]
    errors: list[str] = []
    if len(table.columns) != 11:
        errors.append(f"Expected 11 visible columns; found {len(table.columns)}.")
    rows: list[AuditRow] = []
    for idx, row in enumerate(table.rows[3:], start=4):
        cells = [clean(c.text) for c in row.cells]
        cells += [""] * (11 - len(cells))
        node_id, node_name, input_logic, output_logic, arc_expr, tag, prob, cap, o, ml, p = cells[:11]
        from_node, to_node = parse_arc(arc_expr)
        missing = []
        if not node_id:
            missing.append("Node ID")
        if not node_name:
            missing.append("Node Name")
        if node_name not in {"Final Closeout", "Withdrawn /Rejected", "Withdrawn / Rejected"}:
            for label, val in [
                ("Node Type - Input", input_logic),
                ("Node Type - Output", output_logic),
                ("Arc", arc_expr),
                ("Tag", tag),
                ("Probability", prob),
                ("O", o),
                ("ML", ml),
                ("P", p),
            ]:
                if not val or val == "|":
                    missing.append(label)
        if arc_expr and (not from_node or not to_node):
            missing.append("Explicit From -> To parse")
        status = "Valid" if not missing else "Missing: " + ", ".join(missing)
        if missing:
            errors.append(f"Word row {idx}: {status}.")
        loop_class = "Bounded loop/return" if cap else "Normal/uncapped or terminal"
        pert_class = "Unclassified - blocked until Arc/Tag/O/ML/P complete"
        rows.append(
            AuditRow(
                idx,
                node_id,
                node_name,
                input_logic,
                output_logic,
                arc_expr,
                tag,
                from_node,
                to_node,
                prob,
                cap,
                o,
                ml,
                p,
                loop_class,
                pert_class,
                status,
                f"Table1!A{idx}:K{idx}",
            )
        )
    return rows, errors


def write_rows(path: Path, sheet_name: str, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(55, max(12, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(path)


def build_authoritative_gate(word_path: Path, root: Path) -> dict:
    reports = root / "reports"
    config = root / "config"
    rows, errors = parse_word_panel(word_path)
    checksum = hashlib.sha256(word_path.read_bytes()).hexdigest()
    audit_dicts = [asdict(r) for r in rows]
    write_rows(reports / "01_Input_Audit.xlsx", "Input Audit", audit_dicts)
    node_names = sorted({r.node_name for r in rows if r.node_name})
    arcs = [r for r in rows if r.arc_tag or r.arc_expression]
    loops = [r for r in rows if r.loop_cap]
    network_rows = [
        {"Metric": "Word file", "Value": str(word_path)},
        {"Metric": "File name", "Value": word_path.name},
        {"Metric": "File size", "Value": word_path.stat().st_size},
        {"Metric": "SHA-256", "Value": checksum},
        {"Metric": "Nodes with names", "Value": len(node_names)},
        {"Metric": "Rows", "Value": len(rows)},
        {"Metric": "Rows with Arc Tags", "Value": len([r for r in rows if r.arc_tag])},
        {"Metric": "Rows with explicit Arc expressions", "Value": len([r for r in rows if r.arc_expression])},
        {"Metric": "Loop Cap rows", "Value": len(loops)},
        {"Metric": "Gate Status", "Value": "AUTHORITATIVE INPUT INCOMPLETE - SIMULATION BLOCKED" if errors else "Passed"},
    ]
    for error in errors:
        network_rows.append({"Metric": "Error", "Value": error})
    write_rows(reports / "02_Network_Audit.xlsx", "Network Audit", network_rows)
    write_rows(reports / "04_PERT_Input_Gate.xlsx", "PERT Input Gate", [
        {"Metric": "Stochastic PERT activities extracted", "Value": 0},
        {"Metric": "Expected stochastic PERT activities", "Value": 6},
        {"Metric": "Sum O", "Value": ""},
        {"Metric": "Expected Sum O", "Value": 87},
        {"Metric": "Sum ML", "Value": ""},
        {"Metric": "Expected Sum ML", "Value": 110},
        {"Metric": "Sum P", "Value": ""},
        {"Metric": "Expected Sum P", "Value": 146},
        {"Metric": "Gate Status", "Value": "AUTHORITATIVE INPUT INCOMPLETE - SIMULATION BLOCKED"},
    ])
    config.mkdir(parents=True, exist_ok=True)
    case = {
        "status": "AUTHORITATIVE INPUT INCOMPLETE - SIMULATION BLOCKED" if errors else "parsed",
        "word_file": str(word_path),
        "file_name": word_path.name,
        "file_size": word_path.stat().st_size,
        "sha256": checksum,
        "expected_columns": EXPECTED_COLUMNS,
        "rows": audit_dicts,
        "errors": errors,
    }
    (config / "authoritative_case.json").write_text(json.dumps(case, indent=2), encoding="utf-8")
    return {
        "nodes": len(node_names),
        "arcs": len(arcs),
        "loops": len(loops),
        "pert_activities": 0,
        "sum_o": "",
        "sum_ml": "",
        "sum_p": "",
        "errors": len(errors),
        "status": case["status"],
        "file_name": word_path.name,
        "sha256": checksum,
    }
